#!/usr/bin/env python3
from __future__ import annotations

import argparse
import html as html_lib
import hashlib
import json
import math
import re
import time
import sys
from pathlib import Path
from urllib.parse import quote, urljoin

from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from modules.network_pharmacology.scripts.common import (
    WorkflowError,
    ensure_dir,
    exit_with_errors,
    fail,
    http_get,
    http_get_tcmsp,
    load_params,
    now_utc,
    output_root,
    project_root_from_cwd,
    read_json,
    read_rows,
    unique_preserve,
    write_json,
    write_rows,
)
from modules.network_pharmacology.scripts.herb_source_integrator import (
    fetch_supplemental_herb_sources,
    merge_herb_source_rows,
)
from modules.network_pharmacology.scripts.tcmsp_snapshot import load_snapshot_herb_data


TCMSP_BASE = "https://tcmsp-e.com"
TCMSP_DOWNLOAD_FILES = {
    "herbs": "02_Info_Herbs_Name.xlsx",
    "molecules": "03_Info_Molecules.xlsx",
    "targets": "04_Info_Targets.xlsx",
    "herb_molecules": "23_Herbs_Molecules_Relationships.xlsx",
    "molecule_targets": "34_Molecules_Targets_Relationships.xlsx",
}
TCMSP_DOWNLOAD_URLS = {
    name: f"{TCMSP_BASE}/attachment/tcmspDB/{filename}"
    for name, filename in TCMSP_DOWNLOAD_FILES.items()
}
HERB_ALIASES = {
    "丹参": [
        "丹参",
        "Danshen",
        "Dan Shen",
        "Salvia miltiorrhiza",
        "Radix Salviae",
        "Radix Salviae Miltiorrhizae",
        "Salviae Miltiorrhizae Radix et Rhizoma",
        "Salviae Miltiorrhizae",
    ],
    "苦参": ["苦参", "Kushen", "Ku Shen", "Sophora flavescens", "Sophorae Flavescentis"],
    "赤小豆": [
        "赤小豆",
        "红豆",
        "Chixiaodou",
        "Chi Xiao Dou",
        "Rice Bean",
        "Semen Phaseoli",
        "Vignae Semen",
        "Semen Phaseoli,Vignae Semen",
    ],
}
AMBIGUOUS_HERB_STANDARDIZATION = {
    "红豆": {
        "canonical_herb_name": "赤小豆",
        "alias_source": "curated_ambiguous_common_name",
        "alias_confidence": "medium",
        "match_status": "alias_needs_confirmation",
        "requires_user_confirmation": True,
        "warning": "“红豆”可能对应赤小豆、相思子、红豆蔻等不同药材；本次非交互模式默认按“赤小豆”尝试数据库检索，生物学解释前必须人工确认。",
        "candidates": [
            {
                "name": "赤小豆",
                "reason": "常用中药名，SymMap 2.0 中存在 Rice Bean / Semen Phaseoli,Vignae Semen 条目。",
                "recommended": True,
            },
            {
                "name": "相思子",
                "reason": "部分语境中也被称作红豆，且有毒性风险；不能与赤小豆混用。",
                "recommended": False,
            },
            {
                "name": "红豆蔻",
                "reason": "名称相近但为另一味中药，不应静默合并。",
                "recommended": False,
            },
        ],
    }
}
# Used only when herb search/download endpoints are unavailable. IDs are TCMSP
# molecule IDs that are publicly visible through molecule.php pages.
HERB_MOLECULE_PAGE_SEEDS = {
    "丹参": [
        "MOL000006",
        "MOL007154",
        "MOL007145",
        "MOL007100",
        "MOL007049",
        "MOL007041",
        "MOL007108",
        "MOL007088",
    ],
    "Danshen": [
        "MOL000006",
        "MOL007154",
        "MOL007145",
        "MOL007100",
        "MOL007049",
        "MOL007041",
        "MOL007108",
        "MOL007088",
    ],
    "Salvia miltiorrhiza": [
        "MOL000006",
        "MOL007154",
        "MOL007145",
        "MOL007100",
        "MOL007049",
        "MOL007041",
        "MOL007108",
        "MOL007088",
    ],
}

# Curated mappings for common TCMSP target display names. UniProt fallback is
# still used, but these names are frequent and otherwise hard to resolve from
# prose-like TCMSP labels.
TARGET_NAME_TO_GENE = {
    "Cellular tumor antigen p53": ("TP53", "P04637"),
    "Prostaglandin G/H synthase 1": ("PTGS1", "P23219"),
    "Prostaglandin G/H synthase 2": ("PTGS2", "P35354"),
    "Nitric oxide synthase, inducible": ("NOS2", "P35228"),
    "Nitric-oxide synthase, endothelial": ("NOS3", "P29474"),
    "Tumor necrosis factor": ("TNF", "P01375"),
    "Interleukin-6": ("IL6", "P05231"),
    "Interleukin-1 beta": ("IL1B", "P01584"),
    "Vascular endothelial growth factor A": ("VEGFA", "P15692"),
    "Caspase-3": ("CASP3", "P42574"),
    "Caspase-8": ("CASP8", "Q14790"),
    "Caspase-9": ("CASP9", "P55211"),
    "Apoptosis regulator BAX": ("BAX", "Q07812"),
    "Apoptosis regulator Bcl-2": ("BCL2", "P10415"),
    "Bcl-2-like protein 1": ("BCL2L1", "Q07817"),
    "RAC-alpha serine/threonine-protein kinase": ("AKT1", "P31749"),
    "Mitogen-activated protein kinase 14": ("MAPK14", "Q16539"),
    "Glycogen synthase kinase-3 beta": ("GSK3B", "P49841"),
    "Androgen receptor": ("AR", "P10275"),
    "Estrogen receptor": ("ESR1", "P03372"),
    "Estrogen receptor beta": ("ESR2", "Q92731"),
    "Progesterone receptor": ("PGR", "P06401"),
    "Peroxisome proliferator activated receptor gamma": ("PPARG", "P37231"),
    "Peroxisome proliferator-activated receptor alpha": ("PPARA", "Q07869"),
    "Dipeptidyl peptidase IV": ("DPP4", "P27487"),
    "Heat shock protein HSP 90": ("HSP90AA1", "P07900"),
    "Matrix metalloproteinase-9": ("MMP9", "P14780"),
    "Intercellular adhesion molecule 1": ("ICAM1", "P05362"),
    "Coagulation factor Xa": ("F10", "P00742"),
    "Thrombin": ("F2", "P00734"),
    "Retinoblastoma-associated protein": ("RB1", "P06400"),
    "Retinoic acid receptor RXR-alpha": ("RXRA", "P19793"),
    "Hypoxia-inducible factor 1-alpha": ("HIF1A", "Q16665"),
    "Transcription factor p65": ("RELA", "Q04206"),
    "Transforming growth factor beta-1": ("TGFB1", "P01137"),
    "Insulin receptor": ("INSR", "P06213"),
    "Insulin-like growth factor 1 receptor": ("IGF1R", "P08069"),
    "Carbonic anhydrase II": ("CA2", "P00918"),
    "Alpha-1A adrenergic receptor": ("ADRA1A", "P35348"),
    "Alpha-1B adrenergic receptor": ("ADRA1B", "P35368"),
    "Alpha-1D adrenergic receptor": ("ADRA1D", "P25100"),
    "Beta-2 adrenergic receptor": ("ADRB2", "P07550"),
    "Muscarinic acetylcholine receptor M1": ("CHRM1", "P11229"),
    "Muscarinic acetylcholine receptor M2": ("CHRM2", "P08172"),
    "Muscarinic acetylcholine receptor M3": ("CHRM3", "P20309"),
    "Sodium channel protein type 5 subunit alpha": ("SCN5A", "Q14524"),
    "Potassium voltage-gated channel subfamily H member 2": ("KCNH2", "Q12809"),
    "Glucocorticoid receptor": ("NR3C1", "P04150"),
    "Mineralocorticoid receptor": ("NR3C2", "P08235"),
    "5-hydroxytryptamine receptor 3A": ("HTR3A", "P46098"),
    "Cell division protein kinase 2": ("CDK2", "P24941"),
    "Coagulation factor VII": ("F7", "P08709"),
    "Vascular endothelial growth factor receptor 2": ("KDR", "P35968"),
    "Interleukin-4": ("IL4", "P05112"),
    "NAD-dependent deacetylase sirtuin-1": ("SIRT1", "Q96EB6"),
    "Acetylcholinesterase": ("ACHE", "P22303"),
    "Corticoliberin": ("CRH", "P06850"),
    "Myc proto-oncogene protein": ("MYC", "P01106"),
    "72 kDa type IV collagenase": ("MMP2", "P08253"),
    "CD44 antigen": ("CD44", "P16070"),
    "mRNA of PKA Catalytic Subunit C-alpha": ("PRKACA", "P17612"),
    "Phosphatidylinositol-4,5-bisphosphate 3-kinase catalytic subunit, gamma isoform": ("PIK3CG", "P48736"),
    "Trypsin-1": ("PRSS1", "P07477"),
    "Mu-type opioid receptor": ("OPRM1", "P35372"),
    "Serine/threonine-protein kinase Chk1": ("CHEK1", "O14757"),
    "Proto-oncogene serine/threonine-protein kinase Pim-1": ("PIM1", "P11309"),
    "Cyclin-A2": ("CCNA2", "P20248"),
    "Calcium-activated potassium channel subunit alpha 1": ("KCNMA1", "Q12791"),
    "Nuclear receptor coactivator 1": ("NCOA1", "Q15788"),
    "Nuclear receptor coactivator 2": ("NCOA2", "Q15596"),
    "Gamma-aminobutyric-acid receptor alpha-2 subunit": ("GABRA2", "P47869"),
    "Gamma-aminobutyric acid receptor subunit alpha-1": ("GABRA1", "P14867"),
    "CGMP-inhibited 3',5'-cyclic phosphodiesterase A": ("PDE3A", "Q14432"),
    "Glycogen phosphorylase, muscle form": ("PYGM", "P11217"),
    "Sodium-dependent serotonin transporter": ("SLC6A4", "P31645"),
    "Neuronal acetylcholine receptor protein, alpha-7 chain": ("CHRNA7", "P36544"),
    "Sodium-dependent dopamine transporter": ("SLC6A3", "Q01959"),
    "Amine oxidase [flavin-containing] B": ("MAOB", "P27338"),
    "cAMP-dependent protein kinase inhibitor alpha": ("PKIA", "P61925"),
    "Peroxisome proliferator-activated receptor gamma": ("PPARG", "P37231"),
    "ATP synthase subunit beta, mitochondrial": ("ATP5F1B", "P06576"),
    "NADH-ubiquinone oxidoreductase chain 6": ("MT-ND6", "P03923"),
    "3 beta-hydroxysteroid dehydrogenase/Delta 5-->4-isomerase type 1": ("HSD3B1", "P14060"),
    "3 beta-hydroxysteroid dehydrogenase/Delta 5-->4-isomerase type 2": ("HSD3B2", "P26439"),
    "Aldose reductase": ("AKR1B1", "P15121"),
    "Heparanase": ("HPSE", "Q9Y251"),
    "Epidermal growth factor receptor": ("EGFR", "P00533"),
    "G1/S-specific cyclin-D1": ("CCND1", "P24385"),
    "Cyclin-dependent kinase inhibitor 1": ("CDKN1A", "P38936"),
    "Mitogen-activated protein kinase 1": ("MAPK1", "P28482"),
    "Interleukin-10": ("IL10", "P22301"),
    "Cell division protein kinase 4": ("CDK4", "P11802"),
    "NF-kappa-B inhibitor alpha": ("NFKBIA", "P25963"),
    "Xanthine dehydrogenase/oxidase": ("XDH", "P47989"),
    "DNA topoisomerase 1": ("TOP1", "P11387"),
    "E3 ubiquitin-protein ligase Mdm2": ("MDM2", "Q00987"),
    "Amyloid beta A4 protein": ("APP", "P05067"),
    "Interstitial collagenase": ("MMP1", "P03956"),
    "Proliferating cell nuclear antigen": ("PCNA", "P12004"),
    "Receptor tyrosine-protein kinase erbB-2": ("ERBB2", "P04626"),
    "Heme oxygenase 1": ("HMOX1", "P09601"),
    "Acetyl-CoA carboxylase 1": ("ACACA", "Q13085"),
    "Arachidonate 5-lipoxygenase": ("ALOX5", "P09917"),
    "Aryl hydrocarbon receptor": ("AHR", "P35869"),
    "Baculoviral IAP repeat-containing protein 4": ("XIAP", "P98170"),
    "Baculoviral IAP repeat-containing protein 5": ("BIRC5", "O15392"),
    "C-X-C motif chemokine 2": ("CXCL2", "P19875"),
    "CD40 ligand": ("CD40LG", "P29965"),
    "Caveolin-1": ("CAV1", "Q03135"),
    "Claudin-4": ("CLDN4", "O14493"),
    "Cytochrome P450 3A4": ("CYP3A4", "P08684"),
    "G2/mitotic-specific cyclin-B1": ("CCNB1", "P14635"),
    "Hexokinase-2": ("HK2", "P52789"),
    "Myeloperoxidase": ("MPO", "P05164"),
    "NAD(P)H dehydrogenase [quinone] 1": ("NQO1", "P15559"),
    "NADPH--cytochrome P450 reductase": ("POR", "P16435"),
    "Nuclear receptor subfamily 1 group I member 2": ("NR1I2", "O75469"),
    "Peroxisome proliferator-activated receptor delta": ("PPARD", "Q03181"),
    "Poly [ADP-ribose] polymerase 1": ("PARP1", "P09874"),
    "Prostaglandin E2 receptor EP3 subtype": ("PTGER3", "P43115"),
    "Protein CBFA2T1": ("RUNX1T1", "Q06455"),
    "Runt-related transcription factor 2": ("RUNX2", "Q13950"),
    "Superoxide dismutase [Cu-Zn]": ("SOD1", "P00441"),
    "Type I iodothyronine deiodinase": ("DIO1", "P49895"),
}

CANONICAL_AMBIGUOUS_TARGETS = {
    "Calmodulin": {
        "gene": "CALM1",
        "uniprot": "P0DP23",
        "suggested_gene_symbol": "CALM1; CALM2; CALM3",
        "note": "CALM1/CALM2/CALM3 canonical-selected",
    },
    "DNA topoisomerase II": {
        "gene": "TOP2A",
        "uniprot": "P11388",
        "suggested_gene_symbol": "TOP2A; TOP2B",
        "note": "TOP2A/TOP2B canonical-selected",
    },
}

AMBIGUOUS_TARGETS = {
    "Transcription factor AP-1": {
        "suggested_gene_symbol": "JUN; FOS",
        "note": "AP-1 is a transcription factor complex; not emitted into normalized target network.",
    },
}

IGNORED_TARGET_NAMES = {
    "Beta-lactamase",
    "Ferrichrome-iron receptor",
    "Ig gamma-1 chain C region",
}

COMPOUND_NAME_ALIASES = {
    "补骨脂素": "psoralen",
}


INGREDIENT_FIELDS = [
    "herb_name",
    "ingredient_id",
    "raw_ingredient_name",
    "ingredient_name",
    "normalized_ingredient_name",
    "name_repair_applied",
    "pubchem_cid",
    "smiles",
    "canonical_smiles",
    "isomeric_smiles",
    "inchikey",
    "inchi_key",
    "pubchem_match_status",
    "pubchem_query_name",
    "pubchem_attempted_names",
    "pubchem_error",
    "cid_origin",
    "cid_origin_source_database",
    "pubchem_query_attempted",
    "pubchem_query_status",
    "smiles_origin",
    "inchikey_origin",
    "structure_file_path",
    "structure_status",
    "docking_ready_reason",
    "ob",
    "dl",
    "source",
    "source_list",
    "primary_source",
    "evidence_types",
    "database_ids",
    "activity_filter_basis",
    "ingredient_confidence",
    "source_eligibility",
    "source_eligibility_reason",
    "herbiv_category",
    "network_eligibility",
    "network_eligibility_reason",
    "tcmsp_data_source",
    "passed_filter",
]

LIGAND_FIELDS = [
    "herb_name",
    "ingredient_id",
    "raw_ingredient_name",
    "ingredient_name",
    "normalized_ingredient_name",
    "name_repair_applied",
    "pubchem_cid",
    "smiles",
    "canonical_smiles",
    "isomeric_smiles",
    "inchikey",
    "cid_origin",
    "cid_origin_source_database",
    "pubchem_query_attempted",
    "pubchem_query_status",
    "smiles_origin",
    "inchikey_origin",
    "structure_file_path",
    "structure_status",
    "docking_ready_reason",
    "ob",
    "dl",
    "compound_target_count",
    "source",
    "source_list",
    "primary_source",
    "evidence_types",
    "database_ids",
    "activity_filter_basis",
    "ingredient_confidence",
    "source_eligibility",
    "source_eligibility_reason",
    "herbiv_category",
    "network_eligibility",
    "network_eligibility_reason",
    "tcmsp_data_source",
    "structure_source",
    "pubchem_match_status",
    "pubchem_query_name",
    "pubchem_attempted_names",
    "recommended_for_docking",
]

TARGET_FIELDS = [
    "herb_name",
    "raw_ingredient_name",
    "ingredient_name",
    "normalized_ingredient_name",
    "ingredient_id",
    "pubchem_cid",
    "smiles",
    "canonical_smiles",
    "isomeric_smiles",
    "inchikey",
    "target_gene",
    "uniprot_id",
    "target_name",
    "mapping_status",
    "mapping_note",
    "evidence_type",
    "evidence_class",
    "mapping_validated",
    "score",
    "source",
    "source_list",
    "eligible_source_list",
    "all_source_list",
    "target_evidence_types",
    "source_target_ids",
    "ingredient_confidence",
    "relation_confidence",
    "source_eligibility",
    "source_eligibility_reason",
    "herbiv_category",
    "network_eligibility",
    "network_eligibility_reason",
    "eligibility_status_by_source",
    "eligibility_reason_by_source",
    "tcmsp_data_source",
]

UNMAPPED_TARGET_FIELDS = [
    "target_name",
    "normalized_query",
    "attempted_queries",
    "failure_reason",
    "suggested_gene_symbol",
    "mapping_status",
    "ingredient_ids",
    "ingredient_names",
    "target_row_count",
    "source",
]

AMBIGUOUS_TARGET_FIELDS = UNMAPPED_TARGET_FIELDS

PUBCHEM_UNMATCHED_FIELDS = [
    "ingredient_id",
    "ingredient_name",
    "cleaned_name",
    "attempted_names",
    "reason",
    "used_in_network_analysis",
    "eligible_for_docking",
    "pubchem_error",
    "ob",
    "dl",
    "target_count",
]

PUBCHEM_QUERY_ATTEMPT_FIELDS = [
    "ingredient_id",
    "raw_ingredient_name",
    "ingredient_name",
    "candidate",
    "candidate_type",
    "query_url",
    "http_status",
    "accepted",
    "decision_reason",
    "pubchem_cid",
    "inchikey",
]

MULTI_NAME_FIELDS = [
    "ingredient_id",
    "raw_ingredient_name",
    "ingredient_name",
    "normalized_ingredient_name",
    "source_list",
    "classification",
    "reason",
    "split_names",
]

DEDUPLICATED_ACTIVE_FIELDS = [
    "deduplicated_ingredient_id",
    "identity_key",
    "identity_type",
    "ingredient_id",
    "raw_ingredient_name",
    "ingredient_name",
    "normalized_ingredient_name",
    "name_repair_applied",
    "pubchem_cid",
    "smiles",
    "canonical_smiles",
    "isomeric_smiles",
    "inchikey",
    "source_ingredient_ids",
    "source_ingredient_names",
    "source_herb_names",
    "active_record_count",
    "pubchem_match_status",
    "ob",
    "dl",
    "source",
    "tcmsp_data_source",
]

UNIQUE_LIGAND_FIELDS = [
    "ligand_id",
    "pubchem_cid",
    "inchikey",
    "canonical_smiles",
    "isomeric_smiles",
    "source_ingredient_ids",
    "source_ingredient_names",
    "source_herb_names",
    "ingredient_record_count",
    "compound_target_count",
    "pubchem_match_status",
]


def extract_balanced_array(text: str, start_index: int) -> list[dict]:
    open_index = text.find("[", start_index)
    if open_index == -1:
        fail("Could not find JavaScript array in TCMSP page.")

    depth = 0
    in_string = False
    escape = False
    for index in range(open_index, len(text)):
        char = text[index]
        if in_string:
            if escape:
                escape = False
            elif char == "\\":
                escape = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char == "[":
            depth += 1
        elif char == "]":
            depth -= 1
            if depth == 0:
                raw = text[open_index : index + 1]
                try:
                    return json.loads(raw)
                except json.JSONDecodeError as exc:
                    fail(f"Could not parse TCMSP embedded JSON: {exc}")
    fail("Unclosed JavaScript array in TCMSP page.")


def extract_kendo_grid_data(html: str, grid_id: str) -> list[dict]:
    marker = f'$("#{grid_id}").kendoGrid'
    marker_index = html.find(marker)
    if marker_index == -1:
        fail(f"Could not find TCMSP grid #{grid_id}.")
    data_index = html.find("data:", marker_index)
    if data_index == -1:
        fail(f"Could not find data block for TCMSP grid #{grid_id}.")
    data = extract_balanced_array(html, data_index)
    if not isinstance(data, list):
        fail(f"TCMSP grid #{grid_id} did not contain a list.")
    return data


def extract_html_table_data(html: str, grid_id: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find(id=grid_id)
    if table and table.name != "table":
        table = table.find("table")
    if not table:
        return []

    headers = [cell.get_text(strip=True) for cell in table.find_all("th")]
    rows: list[dict] = []
    for tr in table.find_all("tr"):
        cells = [cell.get_text(strip=True) for cell in tr.find_all(["td", "th"])]
        if not cells or cells == headers:
            continue
        if not headers:
            headers = [f"column_{i + 1}" for i in range(len(cells))]
        rows.append({headers[i] if i < len(headers) else f"column_{i + 1}": value for i, value in enumerate(cells)})
    return rows


def extract_tcmsp_grid_data(html: str, grid_id: str) -> list[dict]:
    try:
        return extract_kendo_grid_data(html, grid_id)
    except WorkflowError as exc:
        rows = extract_html_table_data(html, grid_id)
        if rows:
            return rows
        fail(f"{exc} HTML table fallback for #{grid_id} also returned no rows.")


def validate_tcmsp_herb_html(html: str) -> tuple[bool, str]:
    if not html.strip():
        return False, "empty_response"
    if "Error querying database" in html:
        return False, "tcmsp_detail_unavailable"
    failed = []
    for grid_id in ("grid", "grid2"):
        try:
            rows = extract_tcmsp_grid_data(html, grid_id)
        except WorkflowError as exc:
            failed.append(f"{grid_id}:{exc}")
            continue
        if not rows:
            failed.append(f"{grid_id}:empty")
    if failed:
        return False, "missing_or_unparseable_" + "|".join(failed)
    return True, "valid"


def herb_cache_slug(herb_name: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_\-\u4e00-\u9fff]+", "_", herb_name.strip()).strip("_")
    digest = hashlib.sha1(herb_name.strip().encode("utf-8")).hexdigest()[:8]
    return f"{slug or 'herb'}_{digest}"


def rows_fieldnames(rows: list[dict]) -> list[str]:
    fields: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                fields.append(key)
    return fields


def persistent_herb_cache_dir(project_root: Path, herb_name: str) -> Path:
    return project_root / "data" / "cache" / "tcmsp" / "herbs" / herb_cache_slug(herb_name)


def herb_aliases(herb_name: str) -> list[str]:
    aliases = [herb_name.strip()]
    aliases.extend(HERB_ALIASES.get(herb_name.strip(), []))
    for canonical, values in HERB_ALIASES.items():
        if herb_name.strip().lower() in {value.lower() for value in values}:
            aliases.append(canonical)
            aliases.extend(values)
    return unique_preserve(aliases)


def standardize_herb_name(herb_name: str) -> dict:
    original = herb_name.strip()
    rule = AMBIGUOUS_HERB_STANDARDIZATION.get(original)
    if not rule:
        return {
            "input_herb_name": original,
            "canonical_herb_name": original,
            "standardized_herb_for_lookup": original,
            "standardization_status": "unchanged",
            "alias_source": "",
            "alias_confidence": "",
            "requires_user_confirmation": False,
            "herb_match_status": "",
            "warning": "",
            "standardized_candidates": [],
            "aliases_for_lookup": herb_aliases(original),
        }
    canonical = str(rule["canonical_herb_name"])
    aliases = unique_preserve([original, canonical] + HERB_ALIASES.get(canonical, []))
    return {
        "input_herb_name": original,
        "canonical_herb_name": canonical,
        "standardized_herb_for_lookup": canonical,
        "standardization_status": "ambiguous_alias_defaulted",
        "alias_source": rule.get("alias_source", ""),
        "alias_confidence": rule.get("alias_confidence", ""),
        "requires_user_confirmation": bool(rule.get("requires_user_confirmation", False)),
        "herb_match_status": rule.get("match_status", "alias_needs_confirmation"),
        "warning": rule.get("warning", ""),
        "standardized_candidates": rule.get("candidates", []),
        "aliases_for_lookup": aliases,
    }


def tcmsp_download_cache_dir(project_root: Path) -> Path:
    return project_root / "data" / "cache" / "tcmsp" / "downloads"


def tcmsp_expected_download_paths(project_root: Path) -> list[str]:
    cache_dir = tcmsp_download_cache_dir(project_root)
    return [str(cache_dir / filename) for filename in TCMSP_DOWNLOAD_FILES.values()]


def is_xlsx_file(path: Path) -> bool:
    if not path.exists() or path.stat().st_size < 4:
        return False
    with path.open("rb") as f:
        return f.read(4) == b"PK\x03\x04"


def ensure_tcmsp_downloads(project_root: Path) -> tuple[dict[str, Path], list[str]]:
    cache_dir = ensure_dir(tcmsp_download_cache_dir(project_root))
    paths: dict[str, Path] = {}
    diagnostics: list[str] = []
    for name, filename in TCMSP_DOWNLOAD_FILES.items():
        path = cache_dir / filename
        if is_xlsx_file(path):
            paths[name] = path
            continue
        url = TCMSP_DOWNLOAD_URLS[name]
        try:
            response = http_get_tcmsp(url, timeout=60, retries=1)
            path.write_bytes(response.content)
            if is_xlsx_file(path):
                paths[name] = path
            else:
                preview = response.text[:120].replace("\n", " ") if hasattr(response, "text") else ""
                diagnostics.append(f"{filename}: downloaded content is not xlsx; status={response.status_code}; preview={preview}")
                path.unlink(missing_ok=True)
        except Exception as exc:
            diagnostics.append(f"{filename}: download failed: {type(exc).__name__}: {exc}")
    return paths, diagnostics


def normalize_column_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def read_xlsx_rows(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        fail("openpyxl is required for TCMSP Excel fallback. Run: pip install openpyxl")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_column_name(str(value or "")) for value in rows[0]]
    output = []
    for raw in rows[1:]:
        item = {
            headers[index] if index < len(headers) and headers[index] else f"column_{index + 1}": "" if value is None else str(value).strip()
            for index, value in enumerate(raw)
        }
        if any(item.values()):
            output.append(item)
    return output


def first_value(row: dict[str, str], candidates: list[str]) -> str:
    normalized = {normalize_column_name(key): value for key, value in row.items()}
    for candidate in candidates:
        value = normalized.get(normalize_column_name(candidate), "")
        if value:
            return value
    return ""


def row_contains_any(row: dict[str, str], values: list[str]) -> bool:
    haystack = " | ".join(str(value).lower() for value in row.values())
    return any(value.lower() in haystack for value in values if value)


def molecule_number_from_id(molecule_id: str) -> str:
    value = molecule_id.strip()
    if value.upper().startswith("MOL"):
        return value[3:]
    return value


def molecule_id_from_number(number: str) -> str:
    value = str(number).strip()
    if value.upper().startswith("MOL"):
        return value
    return f"MOL{int(float(value)):06d}" if re.fullmatch(r"\d+(?:\.0+)?", value) else value


def tcmsp_molecule_page_rows(molecule_id: str, herb_name: str) -> tuple[dict, list[dict]]:
    qn = molecule_number_from_id(molecule_id)
    response = http_get_tcmsp(f"{TCMSP_BASE}/molecule.php?qn={quote(qn)}", timeout=45, retries=1)
    if response.status_code != 200:
        fail(f"TCMSP molecule page returned HTTP {response.status_code} for {molecule_id}.")
    mol_rows = extract_js_var_array(response.text, "mol")
    target_rows = extract_js_var_array(response.text, "tar")
    if not mol_rows:
        fail(f"TCMSP molecule page did not expose mol data for {molecule_id}.")
    mol = mol_rows[0]
    mol_id = str(mol.get("MOL_ID", "") or molecule_id)
    ingredient = {
        "MOL_ID": mol_id,
        "molecule_ID": str(mol.get("molecule_ID", "") or molecule_number_from_id(mol_id)),
        "molecule_name": str(mol.get("molecule_name", "")),
        "ob": str(mol.get("ob", "")),
        "dl": str(mol.get("dl", "")),
        "PubChem_ID": str(mol.get("pubchem_cid", "") or mol.get("PubChem_ID", "")),
        "herb_name": herb_name,
    }
    targets = []
    for row in target_rows:
        targets.append(
            {
                "MOL_ID": mol_id,
                "target_name": str(row.get("target_name", "")),
                "SVM_score": str(row.get("SVM_score", "")),
                "RF_score": str(row.get("RF_score", "")),
                "validated": str(row.get("validated", "")),
                "herb_name": herb_name,
            }
        )
    return ingredient, targets


def load_molecule_page_seed(project_root: Path, herb_name: str) -> dict | None:
    seed_ids: list[str] = []
    for alias in herb_aliases(herb_name):
        seed_ids.extend(HERB_MOLECULE_PAGE_SEEDS.get(alias, []))
    seed_ids = unique_preserve(seed_ids)
    if not seed_ids:
        return None

    cache_dir = ensure_dir(persistent_herb_cache_dir(project_root, herb_name))
    ingredient_path = cache_dir / "raw_ingredients.csv"
    target_path = cache_dir / "raw_targets.csv"
    if ingredient_path.exists() and target_path.exists():
        ingredient_rows = read_rows(ingredient_path)
        target_rows = read_rows(target_path)
        if ingredient_rows and target_rows:
            return {
                "source": "tcmsp_molecule_page_seed_cache",
                "official_download_method": "",
                "html": "",
                "ingredient_rows": ingredient_rows,
                "target_rows": target_rows,
                "path": f"{ingredient_path}; {target_path}",
                "download_file_paths": tcmsp_expected_download_paths(project_root),
                "herb_match_field": "curated_seed_alias",
                "matched_herb_id": f"seed:{herb_name}",
                "matched_herb_name": herb_name,
                "is_full_snapshot": False,
                "is_seed_cache": True,
                "source_completeness_level": "seed",
                "completeness_warning": (
                    "TCMSP molecule page seed cache is incomplete and should only be used for workflow smoke testing."
                ),
                "searched_paths": [str(cache_dir)],
            }

    ingredient_rows: list[dict] = []
    target_rows: list[dict] = []
    searched: list[str] = []
    for molecule_id in seed_ids:
        searched.append(f"{TCMSP_BASE}/molecule.php?qn={molecule_number_from_id(molecule_id)}")
        try:
            ingredient, targets = tcmsp_molecule_page_rows(molecule_id, herb_name)
        except Exception as exc:
            searched.append(f"{molecule_id}: {type(exc).__name__}: {exc}")
            continue
        ingredient_rows.append(ingredient)
        target_rows.extend(targets)
        time.sleep(0.1)

    if not ingredient_rows or not target_rows:
        return None
    write_rows(ingredient_path, ingredient_rows, rows_fieldnames(ingredient_rows))
    write_rows(target_path, target_rows, rows_fieldnames(target_rows))
    write_json(
        cache_dir / "cache_manifest.json",
        {
            "herb_name": herb_name,
            "source": "molecule_page_seed",
            "updated_at": now_utc(),
            "ingredient_count": len(ingredient_rows),
            "target_count": len(target_rows),
            "seed_molecule_ids": seed_ids,
        },
    )
    return {
        "source": "tcmsp_molecule_page_seed_cache",
        "official_download_method": "",
        "html": "",
        "ingredient_rows": ingredient_rows,
        "target_rows": target_rows,
        "path": f"{ingredient_path}; {target_path}",
        "download_file_paths": tcmsp_expected_download_paths(project_root),
        "herb_match_field": "curated_seed_alias",
        "matched_herb_id": f"seed:{herb_name}",
        "matched_herb_name": herb_name,
        "is_full_snapshot": False,
        "is_seed_cache": True,
        "source_completeness_level": "seed",
        "completeness_warning": (
            "TCMSP molecule page seed cache is incomplete and should only be used for workflow smoke testing."
        ),
        "searched_paths": searched,
    }


def load_tcmsp_download_cache(project_root: Path, herb_name: str) -> dict | None:
    paths, diagnostics = ensure_tcmsp_downloads(project_root)
    required = {"herbs", "molecules", "targets", "herb_molecules", "molecule_targets"}
    if not required.issubset(paths):
        return {
            "source": "download_cache_unavailable",
            "official_download_method": "xlsx_download",
            "searched_paths": diagnostics + [str(tcmsp_download_cache_dir(project_root))],
            "download_file_paths": [str(path) for path in paths.values()],
        }

    aliases = herb_aliases(herb_name)
    herbs = read_xlsx_rows(paths["herbs"])
    herb_hits = [row for row in herbs if row_contains_any(row, aliases)]
    herb_ids = unique_preserve(
        first_value(row, ["herb_id", "herb_ID", "herb", "HID", "id"])
        for row in herb_hits
    )
    matched_herb_name = ""
    matched_herb_id = ""
    herb_match_field = ""
    if herb_hits:
        matched = herb_hits[0]
        matched_herb_name = first_value(matched, ["herb_name", "name", "chinese_name", "pinyin", "latin_name"]) or next(iter(matched.values()), "")
        matched_herb_id = herb_ids[0] if herb_ids else ""
        herb_match_field = "herbs_name_table"

    herb_molecules = read_xlsx_rows(paths["herb_molecules"])
    related_molecule_ids = []
    for row in herb_molecules:
        if herb_ids and first_value(row, ["herb_id", "herb_ID", "HID", "herb"]).strip() in herb_ids:
            related_molecule_ids.append(first_value(row, ["MOL_ID", "molecule_id", "mol_id", "molecule"]))
        elif row_contains_any(row, aliases):
            related_molecule_ids.append(first_value(row, ["MOL_ID", "molecule_id", "mol_id", "molecule"]))
    related_molecule_ids = unique_preserve(molecule_id_from_number(value) for value in related_molecule_ids if value)
    if not related_molecule_ids:
        return {
            "source": "download_cache_unavailable",
            "official_download_method": "xlsx_parse",
            "searched_paths": [str(path) for path in paths.values()],
            "download_file_paths": [str(path) for path in paths.values()],
            "failure_reason": f"No herb-molecule rows matched aliases: {'; '.join(aliases)}",
        }

    molecules = read_xlsx_rows(paths["molecules"])
    molecule_by_id = {
        molecule_id_from_number(first_value(row, ["MOL_ID", "molecule_id", "mol_id", "molecule"])): row
        for row in molecules
    }
    ingredient_rows = []
    for mol_id in related_molecule_ids:
        row = molecule_by_id.get(mol_id, {})
        ingredient_rows.append(
            {
                "MOL_ID": mol_id,
                "molecule_ID": molecule_number_from_id(mol_id),
                "molecule_name": first_value(row, ["molecule_name", "name", "mol_name"]),
                "ob": first_value(row, ["ob", "oral_bioavailability"]),
                "dl": first_value(row, ["dl", "drug_likeness"]),
                "PubChem_ID": first_value(row, ["pubchem_cid", "pubchem_id", "PubChem_ID", "cid"]),
                "herb_name": herb_name,
            }
        )

    target_info = read_xlsx_rows(paths["targets"])
    target_by_id = {
        first_value(row, ["target_id", "target_ID", "tid", "id"]): first_value(row, ["target_name", "name", "protein_name"])
        for row in target_info
    }
    molecule_targets = read_xlsx_rows(paths["molecule_targets"])
    target_rows = []
    related_set = set(related_molecule_ids)
    for row in molecule_targets:
        mol_id = molecule_id_from_number(first_value(row, ["MOL_ID", "molecule_id", "mol_id", "molecule"]))
        if mol_id not in related_set:
            continue
        target_name = first_value(row, ["target_name", "name", "protein_name"])
        if not target_name:
            target_name = target_by_id.get(first_value(row, ["target_id", "target_ID", "tid"]), "")
        target_rows.append(
            {
                "MOL_ID": mol_id,
                "target_name": target_name,
                "SVM_score": first_value(row, ["SVM_score", "svm", "score"]),
                "RF_score": first_value(row, ["RF_score", "rf"]),
                "validated": first_value(row, ["validated"]),
                "herb_name": herb_name,
            }
        )

    if not ingredient_rows or not target_rows:
        return None
    return {
        "source": "tcmsp_official_download_files",
        "official_download_method": "xlsx_download_cache",
        "html": "",
        "ingredient_rows": ingredient_rows,
        "target_rows": target_rows,
        "path": "; ".join(str(path) for path in paths.values()),
        "download_file_paths": [str(path) for path in paths.values()],
        "herb_match_field": herb_match_field,
        "matched_herb_id": matched_herb_id,
        "matched_herb_name": matched_herb_name or herb_name,
        "is_full_snapshot": False,
        "is_seed_cache": False,
        "is_official_download_files": True,
        "source_completeness_level": "official_download_files",
        "completeness_warning": "",
        "searched_paths": [str(path) for path in paths.values()],
    }


def write_herb_persistent_cache(
    *,
    project_root: Path,
    herb_name: str,
    detail_url: str,
    html: str,
    ingredient_rows: list[dict],
    target_rows: list[dict],
    source: str,
) -> Path:
    cache_dir = ensure_dir(persistent_herb_cache_dir(project_root, herb_name))
    (cache_dir / "latest_herb_detail.html").write_text(html, encoding="utf-8")
    write_rows(cache_dir / "raw_ingredients.csv", ingredient_rows, rows_fieldnames(ingredient_rows))
    write_rows(cache_dir / "raw_targets.csv", target_rows, rows_fieldnames(target_rows))
    write_json(
        cache_dir / "cache_manifest.json",
        {
            "herb_name": herb_name,
            "detail_url": detail_url,
            "source": source,
            "updated_at": now_utc(),
            "ingredient_count": len(ingredient_rows),
            "target_count": len(target_rows),
            "files": [
                str(cache_dir / "latest_herb_detail.html"),
                str(cache_dir / "raw_ingredients.csv"),
                str(cache_dir / "raw_targets.csv"),
            ],
        },
    )
    return cache_dir


def load_persistent_herb_cache(project_root: Path, herb_name: str) -> dict | None:
    cache_dir = persistent_herb_cache_dir(project_root, herb_name)
    searched = [str(cache_dir)]
    html_path = cache_dir / "latest_herb_detail.html"
    if html_path.exists():
        html = html_path.read_text(encoding="utf-8", errors="replace")
        valid, reason = validate_tcmsp_herb_html(html)
        if valid:
            return {
                "source": "tcmsp_persistent_cache",
                "html": html,
                "ingredient_rows": extract_tcmsp_grid_data(html, "grid"),
                "target_rows": extract_tcmsp_grid_data(html, "grid2"),
                "path": str(html_path),
                "is_full_snapshot": False,
                "is_seed_cache": False,
                "source_completeness_level": "persistent_cache",
                "completeness_warning": "Using persistent TCMSP herb cache because live data was unavailable.",
                "searched_paths": searched,
            }
        searched.append(f"{html_path}: {reason}")
    ingredient_path = cache_dir / "raw_ingredients.csv"
    target_path = cache_dir / "raw_targets.csv"
    if ingredient_path.exists() and target_path.exists():
        ingredient_rows = read_rows(ingredient_path)
        target_rows = read_rows(target_path)
        if ingredient_rows and target_rows:
            return {
                "source": "tcmsp_persistent_cache",
                "html": "",
                "ingredient_rows": ingredient_rows,
                "target_rows": target_rows,
                "path": f"{ingredient_path}; {target_path}",
                "is_full_snapshot": False,
                "is_seed_cache": False,
                "source_completeness_level": "persistent_cache",
                "completeness_warning": "Using persistent TCMSP herb cache because live data was unavailable.",
                "searched_paths": searched,
            }
    return None


def historical_run_matches_herb(run_dir: Path, herb_name: str) -> bool:
    for rel in ("00_input/input_summary.json", "00_input/params.resolved.yaml"):
        path = run_dir / rel
        if not path.exists():
            continue
        try:
            if path.suffix == ".json":
                data = read_json(path)
            else:
                import yaml

                with path.open("r", encoding="utf-8") as f:
                    data = yaml.safe_load(f)
            if isinstance(data, dict):
                candidate = data.get("herb_name") or data.get("input", {}).get("herb_name")
                if candidate == herb_name:
                    return True
        except Exception:
            continue
    return False


def load_historical_herb_cache(project_root: Path, herb_name: str, current_out_root: Path) -> dict | None:
    runs_dir = project_root / "results" / "network_pharmacology_runs"
    if not runs_dir.exists():
        return None
    candidates = sorted(
        [path for path in runs_dir.iterdir() if path.is_dir() and path != current_out_root],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    searched: list[str] = []
    for run_dir in candidates:
        if not historical_run_matches_herb(run_dir, herb_name):
            continue
        html_path = run_dir / "_cache" / "tcmsp_herb_detail.html"
        searched.append(str(html_path))
        if not html_path.exists():
            continue
        html = html_path.read_text(encoding="utf-8", errors="replace")
        valid, reason = validate_tcmsp_herb_html(html)
        if not valid:
            searched.append(f"{html_path}: {reason}")
            continue
        return {
            "source": "tcmsp_historical_run_cache",
            "html": html,
            "ingredient_rows": extract_tcmsp_grid_data(html, "grid"),
            "target_rows": extract_tcmsp_grid_data(html, "grid2"),
            "path": str(html_path),
            "is_full_snapshot": False,
            "is_seed_cache": False,
            "source_completeness_level": "historical_run_cache",
            "completeness_warning": "Using a previous run cache because live TCMSP data was unavailable.",
            "searched_paths": searched,
        }
    return None


def contains_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def tcmsp_token() -> str:
    return "fae1b3f834b798901e77ee835d7e468f"


def default_herb_detail_url(herb_name: str, qsr: str | None = None) -> str:
    field = qsr or ("herb_cn_name" if contains_chinese(herb_name) else "herb_en_name")
    return f"{TCMSP_BASE}/tcmspsearch.php?qr={quote(herb_name)}&qsr={field}&token={tcmsp_token()}"


def response_debug_meta(response, *, parser_stage: str, failure_reason: str = "") -> dict:
    text = response.text or ""
    return {
        "request_url": response.request.url if response.request else "",
        "final_url": response.url,
        "status_code": response.status_code,
        "response_length": len(text),
        "content_type": response.headers.get("content-type", ""),
        "first_500_chars": text[:500],
        "parser_stage": parser_stage,
        "failure_reason": failure_reason,
    }


def grid_required_fields_valid(rows: list[dict], required_any: list[str]) -> bool:
    if not rows:
        return False
    available = {str(key).lower() for row in rows[:5] for key in row}
    return any(field.lower() in available for field in required_any)


def classify_tcmsp_herb_page(html: str, status_code: int | str = "") -> dict:
    text = html or ""
    grid_found = ('$("#grid").kendoGrid' in text) or bool(BeautifulSoup(text, "html.parser").find(id="grid"))
    grid2_found = ('$("#grid2").kendoGrid' in text) or bool(BeautifulSoup(text, "html.parser").find(id="grid2"))
    if not text.strip():
        page_type = "empty_response"
        reason = "empty_response"
    elif "Error querying database" in text:
        page_type = "tcmsp_database_error_page"
        reason = "tcmsp_detail_unavailable"
    elif grid_found or grid2_found:
        page_type = "herb_detail_candidate"
        reason = ""
    elif "TCMSP - Traditional Chinese Medicine Systems Pharmacology Database" in text:
        page_type = "tcmsp_html_without_grid"
        reason = "tcmsp_html_without_grid"
    else:
        page_type = "unknown_html"
        reason = "unknown_html_without_grid"

    grid_row_count = 0
    grid2_row_count = 0
    grid_required_valid = False
    grid2_required_valid = False
    parse_success = False
    parse_reason = reason
    if grid_found and grid2_found:
        try:
            grid_rows = extract_tcmsp_grid_data(text, "grid")
            grid2_rows = extract_tcmsp_grid_data(text, "grid2")
            grid_row_count = len(grid_rows)
            grid2_row_count = len(grid2_rows)
            grid_required_valid = grid_required_fields_valid(grid_rows, ["MOL_ID", "molecule_name", "molecule_ID", "ob", "dl"])
            grid2_required_valid = grid_required_fields_valid(grid2_rows, ["MOL_ID", "target_name", "target_gene"])
            parse_success = bool(grid_rows and grid2_rows and grid_required_valid and grid2_required_valid)
            if not grid_rows or not grid2_rows:
                parse_reason = "grid_or_grid2_empty"
            elif not grid_required_valid:
                parse_reason = "grid_missing_required_ingredient_fields"
            elif not grid2_required_valid:
                parse_reason = "grid2_missing_required_target_fields"
        except WorkflowError as exc:
            parse_reason = str(exc)
    elif not parse_reason:
        parse_reason = "missing_grid_or_grid2"

    return {
        "remote_http_ok": status_code == 200 or str(status_code) == "200",
        "remote_html_valid": bool(text.strip()) and "Error querying database" not in text,
        "remote_grid_found": grid_found,
        "remote_grid2_found": grid2_found,
        "remote_grid_row_count": grid_row_count,
        "remote_grid2_row_count": grid2_row_count,
        "remote_grid_required_fields_valid": grid_required_valid,
        "remote_grid2_required_fields_valid": grid2_required_valid,
        "remote_parse_success": parse_success,
        "remote_page_type": page_type,
        "failure_reason": "" if parse_success else parse_reason,
    }


def remote_field_order(query: str) -> list[str]:
    if contains_chinese(query):
        return [
            "herb_cn_name",
            "herb_all_name",
            "herb_pinyin",
            "herb_pinyin_name",
            "herb_name_pinyin",
            "herb_latin_name",
            "herb_en_name",
        ]
    return [
        "herb_en_name",
        "herb_latin_name",
        "herb_pinyin",
        "herb_pinyin_name",
        "herb_name_pinyin",
        "herb_all_name",
        "herb_cn_name",
    ]


def build_remote_search_attempt_specs(herb_name: str) -> list[dict]:
    specs: list[dict] = []
    aliases = herb_aliases(herb_name)
    for alias in aliases:
        search_field = "herb_cn_name" if contains_chinese(alias) else "herb_en_name"
        specs.append(
            {
                "attempt_type": "search",
                "query": alias,
                "field": search_field,
                "param_name": "qs",
                "url": f"{TCMSP_BASE}/tcmspsearch.php?qs={search_field}&q={quote(alias)}&token={tcmsp_token()}",
            }
        )
        specs.append(
            {
                "attempt_type": "search",
                "query": alias,
                "field": "herb_all_name",
                "param_name": "qs",
                "url": f"{TCMSP_BASE}/tcmspsearch.php?qs=herb_all_name&q={quote(alias)}&token={tcmsp_token()}",
            }
        )
    for alias in aliases:
        for field in remote_field_order(alias):
            specs.append(
                {
                    "attempt_type": "detail",
                    "query": alias,
                    "field": field,
                    "param_name": "qsr",
                    "url": default_herb_detail_url(alias, field),
                }
            )
    output = []
    seen = set()
    for spec in specs:
        key = (spec["attempt_type"], spec["field"], spec["query"], spec["url"])
        if key in seen:
            continue
        seen.add(key)
        output.append(spec)
    return output


def parse_search_candidate_urls(html: str) -> list[str]:
    urls: list[str] = []
    soup = BeautifulSoup(html, "html.parser")
    for link in soup.find_all("a", href=True):
        href = str(link["href"])
        if "tcmspsearch.php" in href and "qsr=" in href:
            urls.append(urljoin(TCMSP_BASE + "/", href))
    for match in re.finditer(r"tcmspsearch\.php\?qr=([^\"'<>]+?)&qsr=([^\"'<>]+?)&token=([a-f0-9]+)", html):
        urls.append(urljoin(TCMSP_BASE + "/", match.group(0)))
    return unique_preserve(urls)


def write_tcmsp_debug_meta(out_root: Path, diagnostic: dict) -> None:
    debug_dir = ensure_dir(out_root / "_debug")
    write_json(debug_dir / "tcmsp_fetch_meta.json", diagnostic)


def find_herb_detail_url(herb_name: str, debug_dir: Path | None = None) -> tuple[str, dict]:
    specs = build_remote_search_attempt_specs(herb_name)
    fallback_url = default_herb_detail_url(herb_name)
    primary_url = specs[0]["url"] if specs else fallback_url
    meta = {
        "request_url": "",
        "final_url": "",
        "status_code": "",
        "response_length": 0,
        "content_type": "",
        "first_500_chars": "",
        "parser_stage": "remote_search_planning",
        "failure_reason": "no_remote_attempt_succeeded",
        "resolved_detail_url": fallback_url,
        "remote_primary_search_url": primary_url,
        "remote_search_url": primary_url,
        "remote_last_attempt_url": "",
        "remote_candidate_count": 0,
        "remote_candidates": [],
        "remote_selected_candidate_url": "",
        "remote_selected_candidate_text": "",
        "remote_selected_candidate_match_field": "",
        "remote_selected_candidate_match_score": "",
        "remote_candidate_failure_reason": "not_attempted",
        "remote_search_attempts": [],
    }
    if debug_dir:
        ensure_dir(debug_dir)
    consecutive_database_errors = 0
    for index, spec in enumerate(specs, start=1):
        attempt = {
            "attempt_index": index,
            "attempt_type": spec["attempt_type"],
            "query": spec["query"],
            "field": spec["field"],
            "qsr": spec["field"] if spec.get("param_name") == "qsr" else "",
            "qs": spec["field"] if spec.get("param_name") == "qs" else "",
            "url": spec["url"],
            "selected": False,
            "candidate_count": 0,
        }
        meta["remote_last_attempt_url"] = spec["url"]
        try:
            response = http_get_tcmsp(spec["url"], timeout=12, retries=1)
            page_meta = response_debug_meta(response, parser_stage=f"{spec['attempt_type']}_attempt")
            classification = classify_tcmsp_herb_page(response.text, response.status_code)
            if classification.get("remote_page_type") == "tcmsp_database_error_page":
                consecutive_database_errors += 1
            else:
                consecutive_database_errors = 0
            attempt.update(page_meta)
            attempt.update(classification)
            attempt["http_status"] = response.status_code
            if debug_dir:
                filename = f"tcmsp_{spec['attempt_type']}_attempt_{index:02d}.html"
                (debug_dir / filename).write_text(response.text, encoding="utf-8")
                if index == 1 and spec["attempt_type"] == "search":
                    (debug_dir / "tcmsp_search_page.html").write_text(response.text, encoding="utf-8")
            if spec["attempt_type"] == "search":
                candidates = parse_search_candidate_urls(response.text)
                attempt["candidate_detail_urls"] = candidates[:20]
                attempt["candidate_count"] = len(candidates)
                meta["remote_candidate_count"] = max(int(meta.get("remote_candidate_count") or 0), len(candidates))
                if candidates:
                    meta["remote_candidates"] = candidates[:20]
                if candidates:
                    attempt["selected"] = True
                    meta["remote_search_attempts"].append(attempt)
                    meta["parser_stage"] = "search_candidate_link"
                    meta["resolved_detail_url"] = candidates[0]
                    meta["remote_selected_candidate_url"] = candidates[0]
                    meta["remote_selected_candidate_text"] = candidates[0]
                    meta["remote_selected_candidate_match_field"] = spec["field"]
                    meta["remote_selected_candidate_match_score"] = "1.0"
                    meta["remote_candidate_failure_reason"] = ""
                    meta.update(page_meta)
                    meta.update(classification)
                    return candidates[0], meta
                if classification.get("remote_page_type") == "tcmsp_database_error_page":
                    attempt["candidate_failure_reason"] = "search_page_database_error"
                    meta["remote_candidate_failure_reason"] = "search_page_database_error"
                else:
                    attempt["candidate_failure_reason"] = "candidate_link_not_found"
                    meta["remote_candidate_failure_reason"] = "candidate_link_not_found"
            elif classification["remote_parse_success"]:
                attempt["selected"] = True
                meta["remote_search_attempts"].append(attempt)
                meta["parser_stage"] = "detail_attempt_valid"
                meta["resolved_detail_url"] = spec["url"]
                meta["remote_selected_candidate_url"] = spec["url"]
                meta["remote_selected_candidate_text"] = spec["query"]
                meta["remote_selected_candidate_match_field"] = spec["field"]
                meta["remote_selected_candidate_match_score"] = "1.0"
                meta["remote_candidate_failure_reason"] = ""
                meta.update(page_meta)
                meta.update(classification)
                return spec["url"], meta
        except Exception as exc:
            attempt.update(
                {
                    "status_code": "",
                    "response_length": 0,
                    "content_type": "",
                    "first_500_chars": "",
                    "remote_http_ok": False,
                    "remote_html_valid": False,
                    "remote_grid_found": False,
                    "remote_grid2_found": False,
                    "remote_parse_success": False,
                    "remote_page_type": "request_error",
                    "failure_reason": f"{type(exc).__name__}: {exc}",
                    "candidate_failure_reason": "request_error",
                }
            )
        meta["remote_search_attempts"].append(attempt)
        if consecutive_database_errors >= 3:
            meta["failure_reason"] = "tcmsp_database_error_page"
            meta["remote_candidate_failure_reason"] = "repeated_tcmsp_database_error_page"
            break

    if meta["remote_search_attempts"]:
        last = meta["remote_search_attempts"][-1]
        meta.update({key: last.get(key, meta.get(key, "")) for key in (
            "request_url",
            "final_url",
            "status_code",
            "response_length",
            "content_type",
            "first_500_chars",
            "remote_http_ok",
            "remote_html_valid",
            "remote_grid_found",
            "remote_grid2_found",
            "remote_grid_row_count",
            "remote_grid2_row_count",
            "remote_grid_required_fields_valid",
            "remote_grid2_required_fields_valid",
            "remote_parse_success",
            "remote_page_type",
            "failure_reason",
        )})
    return fallback_url, meta


def write_minimal_failure_report(*, params: dict, out_root: Path, failed_step: str, error: str, diagnostic: dict) -> None:
    report_dir = ensure_dir(out_root / "08_report")
    debug_files = [
        str(out_root / "_debug" / "tcmsp_search_page.html"),
        str(out_root / "_debug" / "tcmsp_detail_page.html"),
        str(out_root / "_debug" / "tcmsp_fetch_meta.json"),
        str(out_root / "01_ingredients" / "tcmsp_fetch_diagnostic.json"),
    ]
    qc = {
        "module": "network_pharmacology",
        "status": "failed",
        "failed_step": failed_step,
        "generated_at": now_utc(),
        "project_name": params.get("project", {}).get("name", ""),
        "input_mode": params.get("input", {}).get("mode", "herb"),
        "herb_name": params.get("input", {}).get("herb_name", ""),
        "compound_name": params.get("input", {}).get("compound_name", ""),
        "disease_name": params.get("input", {}).get("disease_name", ""),
        "can_prioritize_targets": False,
        "can_prepare_targets": False,
        "can_prepare_ligands": False,
        "can_enter_docking": False,
        "tcmsp_data_source": diagnostic.get("tcmsp_data_source", "unavailable"),
        "tcmsp_fallback_used": bool(diagnostic.get("tcmsp_fallback_used", False)),
        "error": error,
        "debug_file_paths": debug_files,
        "warnings": [error],
        "counts": {
            "active_ingredients": 0,
            "intersection_targets": 0,
            "docking_targets": 0,
            "docking_ligands": 0,
        },
    }
    md = f"""# Network Pharmacology Report

Generated at: {qc["generated_at"]}

## Status

- Status: failed
- Failed step: {failed_step}
- Can enter docking: false

## Input

- Input mode: {qc["input_mode"]}
- Herb/formula: {qc["herb_name"]}
- Compound: {qc["compound_name"]}
- Disease: {qc["disease_name"]}

## Error

{error}

## TCMSP Debug Files

""" + "\n".join(f"- `{path}`" for path in debug_files) + "\n"
    html_body = html_lib.escape(md).replace("\n", "<br>\n")
    html_report = f"<!doctype html><html><head><meta charset=\"utf-8\"><title>NPFlow Failed Report</title></head><body><main><pre>{html_body}</pre></main></body></html>"
    write_json(report_dir / "qc_report.json", qc)
    (report_dir / "network_pharmacology_report.md").write_text(md, encoding="utf-8")
    (report_dir / "network_pharmacology_report.html").write_text(html_report, encoding="utf-8")


def fetch_herb_detail_data(
    *,
    project_root: Path,
    out_root: Path,
    cache_dir: Path,
    herb_name: str,
) -> tuple[list[dict], list[dict], str, str, dict]:
    debug_dir = ensure_dir(out_root / "_debug")
    remote_html_path = cache_dir / "tcmsp_remote_response.html"
    selected_html_path = cache_dir / "tcmsp_herb_detail.html"
    debug_detail_html_path = debug_dir / "tcmsp_detail_page.html"
    detail_url, search_meta = find_herb_detail_url(herb_name, debug_dir)
    diagnostic = {
        "requested_herb": herb_name,
        "input_herb_name": herb_name,
        "resolved_herb_name": search_meta.get("resolved_herb_name", ""),
        "herb_en_name": search_meta.get("herb_en_name", ""),
        "detail_url": detail_url,
        "remote_primary_search_url": search_meta.get("remote_primary_search_url", ""),
        "remote_search_url": search_meta.get("remote_search_url", search_meta.get("remote_primary_search_url", "")),
        "remote_last_attempt_url": search_meta.get("remote_last_attempt_url", ""),
        "remote_selected_candidate_url": search_meta.get("remote_selected_candidate_url", ""),
        "remote_selected_candidate_text": search_meta.get("remote_selected_candidate_text", ""),
        "remote_selected_candidate_match_field": search_meta.get("remote_selected_candidate_match_field", ""),
        "remote_selected_candidate_match_score": search_meta.get("remote_selected_candidate_match_score", ""),
        "remote_candidate_count": search_meta.get("remote_candidate_count", 0),
        "remote_candidates": search_meta.get("remote_candidates", []),
        "remote_candidate_failure_reason": search_meta.get("remote_candidate_failure_reason", ""),
        "request_url": detail_url,
        "final_url": "",
        "http_status": "",
        "status_code": "",
        "response_length": 0,
        "content_type": "",
        "first_500_chars": "",
        "parser_stage": "detail_request",
        "detected_reason": "",
        "failure_reason": "",
        "fallback_searched_paths": [],
        "final_decision": "",
        "tcmsp_data_source": "live_tcmsp_remote",
        "tcmsp_fallback_used": False,
        "selected_cache_path": "",
        "fallback_used": False,
        "cache_path": "",
        "search_meta": search_meta,
        "remote_search_attempts": search_meta.get("remote_search_attempts", []),
        "remote_attempted": True,
        "remote_success": False,
        "remote_failure_reason": "",
        "remote_detail_url": detail_url,
        "remote_http_status": "",
        "remote_response_length": 0,
        "remote_http_ok": bool(search_meta.get("remote_http_ok", False)),
        "remote_html_valid": bool(search_meta.get("remote_html_valid", False)),
        "remote_grid_found": bool(search_meta.get("remote_grid_found", False)),
        "remote_grid2_found": bool(search_meta.get("remote_grid2_found", False)),
        "remote_grid_row_count": int(search_meta.get("remote_grid_row_count", 0) or 0),
        "remote_grid2_row_count": int(search_meta.get("remote_grid2_row_count", 0) or 0),
        "remote_grid_required_fields_valid": bool(search_meta.get("remote_grid_required_fields_valid", False)),
        "remote_grid2_required_fields_valid": bool(search_meta.get("remote_grid2_required_fields_valid", False)),
        "remote_parse_success": bool(search_meta.get("remote_parse_success", False)),
        "remote_page_type": search_meta.get("remote_page_type", ""),
        "live_tcmsp_scrape_validated": False,
        "fallback_snapshot_validated": False,
        "final_tcmsp_source_is_live": False,
        "actual_source_attempt_order": ["live_tcmsp_remote"],
        "fallback_after_remote": False,
        "generated_at": now_utc(),
    }

    skip_detail_request = (
        search_meta.get("remote_candidate_failure_reason") == "repeated_tcmsp_database_error_page"
        and not search_meta.get("remote_selected_candidate_url")
    )
    if skip_detail_request:
        html = ""
        reason = "tcmsp_database_error_page"
        valid = False
        diagnostic.update(
            {
                "parser_stage": "detail_request_skipped_after_repeated_database_errors",
                "remote_page_type": "tcmsp_database_error_page",
                "remote_http_ok": False,
                "remote_html_valid": False,
                "remote_grid_found": False,
                "remote_grid2_found": False,
                "remote_parse_success": False,
            }
        )
    else:
        response = http_get_tcmsp(detail_url, timeout=25, retries=1)
        diagnostic.update(response_debug_meta(response, parser_stage="detail_page"))
        diagnostic["http_status"] = response.status_code
        diagnostic["status_code"] = response.status_code
        diagnostic["remote_http_status"] = response.status_code
        diagnostic["remote_response_length"] = len(response.text or "")
        diagnostic["remote_last_attempt_url"] = detail_url
        html = response.text
        remote_html_path.write_text(html, encoding="utf-8")
        debug_detail_html_path.write_text(html, encoding="utf-8")
        classification = classify_tcmsp_herb_page(html, response.status_code)
        diagnostic.update(classification)
        valid, reason = validate_tcmsp_herb_html(html)
    diagnostic["detected_reason"] = reason
    diagnostic["failure_reason"] = "" if valid else reason
    diagnostic["remote_failure_reason"] = "" if valid else reason
    if not skip_detail_request and response.status_code == 200 and valid:
        ingredient_rows = extract_tcmsp_grid_data(html, "grid")
        target_rows = extract_tcmsp_grid_data(html, "grid2")
        selected_html_path.write_text(html, encoding="utf-8")
        cache_path = write_herb_persistent_cache(
            project_root=project_root,
            herb_name=herb_name,
            detail_url=detail_url,
            html=html,
            ingredient_rows=ingredient_rows,
            target_rows=target_rows,
            source="remote",
        )
        diagnostic["final_decision"] = "used_remote_tcmsp_detail"
        diagnostic["tcmsp_source_conclusion"] = "Live TCMSP scrape succeeded and was used."
        diagnostic["tcmsp_data_source"] = "live_tcmsp_remote"
        diagnostic["remote_success"] = True
        diagnostic["remote_parse_success"] = True
        diagnostic["remote_grid_row_count"] = len(ingredient_rows)
        diagnostic["remote_grid2_row_count"] = len(target_rows)
        diagnostic["remote_grid_required_fields_valid"] = True
        diagnostic["remote_grid2_required_fields_valid"] = True
        diagnostic["live_tcmsp_scrape_validated"] = True
        diagnostic["final_tcmsp_source_is_live"] = True
        diagnostic["selected_cache_path"] = str(cache_path)
        diagnostic["cache_path"] = str(cache_path)
        diagnostic["is_full_snapshot"] = False
        diagnostic["is_seed_cache"] = False
        diagnostic["source_completeness_level"] = "remote_herb_page"
        diagnostic["raw_ingredient_count"] = len(ingredient_rows)
        diagnostic["raw_target_relation_count"] = len(target_rows)
        write_tcmsp_debug_meta(out_root, diagnostic)
        return ingredient_rows, target_rows, detail_url, html, diagnostic

    searched_paths: list[str] = []
    fallback = None
    snapshot_fallback = load_snapshot_herb_data(
        project_root,
        herb_name,
        aliases=herb_aliases(herb_name),
        seed_cache_dir=persistent_herb_cache_dir(project_root, herb_name),
    )
    diagnostic["actual_source_attempt_order"].append("tcmsp_full_snapshot")
    if snapshot_fallback:
        searched_paths.extend(snapshot_fallback.get("searched_paths", []))
        fallback = snapshot_fallback
    download_fallback = None
    if not fallback:
        diagnostic["actual_source_attempt_order"].append("tcmsp_official_download_files")
        download_fallback = load_tcmsp_download_cache(project_root, herb_name)
    if download_fallback:
        searched_paths.extend(download_fallback.get("searched_paths", []))
        if download_fallback.get("source") == "tcmsp_official_download_files":
            fallback = download_fallback
    if not fallback:
        diagnostic["actual_source_attempt_order"].append("molecule_page_seed_cache")
        seed_fallback = load_molecule_page_seed(project_root, herb_name)
        if seed_fallback:
            searched_paths.extend(seed_fallback.get("searched_paths", []))
            fallback = seed_fallback
    if not fallback:
        diagnostic["actual_source_attempt_order"].append("persistent_cache")
        persistent_dir = persistent_herb_cache_dir(project_root, herb_name)
        searched_paths.append(str(persistent_dir))
        fallback = load_persistent_herb_cache(project_root, herb_name)
        if fallback:
            searched_paths.extend(fallback.get("searched_paths", []))
    if not fallback:
        diagnostic["actual_source_attempt_order"].append("historical_run_cache")
        fallback = load_historical_herb_cache(project_root, herb_name, out_root)
        if fallback:
            searched_paths.extend(fallback.get("searched_paths", []))

    if fallback:
        diagnostic["fallback_searched_paths"] = searched_paths
        diagnostic["final_decision"] = f"used_{fallback['source']}"
        diagnostic["tcmsp_data_source"] = fallback["source"]
        if bool(fallback.get("is_full_snapshot", False)):
            if diagnostic.get("remote_page_type") == "tcmsp_database_error_page":
                diagnostic["tcmsp_source_conclusion"] = (
                    "TCMSP live page returned a database error and no valid grid/grid2; "
                    "local full snapshot was used."
                )
            else:
                diagnostic["tcmsp_source_conclusion"] = "Live TCMSP scrape failed or was unparseable; local full snapshot was used."
        else:
            diagnostic["tcmsp_source_conclusion"] = f"Live TCMSP scrape failed or was unparseable; fallback source {fallback['source']} was used."
        diagnostic["tcmsp_fallback_used"] = True
        diagnostic["fallback_after_remote"] = True
        diagnostic["selected_cache_path"] = fallback.get("path", "")
        diagnostic["fallback_used"] = True
        diagnostic["cache_path"] = fallback.get("path", "")
        diagnostic["official_download_method"] = fallback.get("official_download_method", "")
        diagnostic["download_file_paths"] = fallback.get("download_file_paths", [])
        diagnostic["source_urls"] = fallback.get("source_urls", [])
        diagnostic["source_detail"] = fallback.get("source_detail", "")
        diagnostic["snapshot_manifest_path"] = fallback.get("snapshot_manifest_path", "")
        diagnostic["snapshot_db_path"] = fallback.get("snapshot_db_path", "")
        diagnostic["snapshot_source_type"] = fallback.get("snapshot_source_type", "")
        diagnostic["snapshot_row_counts"] = fallback.get("snapshot_row_counts", {})
        diagnostic["snapshot_audit_path"] = fallback.get("snapshot_audit_path", "")
        diagnostic["is_full_snapshot"] = bool(fallback.get("is_full_snapshot", False))
        diagnostic["is_seed_cache"] = bool(fallback.get("is_seed_cache", False))
        diagnostic["is_official_download_files"] = bool(fallback.get("is_official_download_files", False))
        diagnostic["fallback_snapshot_validated"] = bool(fallback.get("is_full_snapshot", False))
        diagnostic["final_tcmsp_source_is_live"] = False
        diagnostic["source_completeness_level"] = fallback.get("source_completeness_level", "")
        diagnostic["completeness_warning"] = fallback.get("completeness_warning", "")
        diagnostic["herb_match_field"] = fallback.get("herb_match_field", "")
        diagnostic["herb_match_score"] = fallback.get("herb_match_score", "")
        diagnostic["input_herb_name"] = fallback.get("input_herb_name", herb_name)
        diagnostic["herb_match_type"] = fallback.get("herb_match_type", "")
        diagnostic["herb_exact_match"] = bool(fallback.get("herb_exact_match", False))
        diagnostic["herb_alias_match"] = bool(fallback.get("herb_alias_match", False))
        diagnostic["herb_fuzzy_match_used"] = bool(fallback.get("herb_fuzzy_match_used", False))
        diagnostic["herb_match_requires_user_confirmation"] = bool(
            fallback.get("herb_match_requires_user_confirmation", False)
        )
        diagnostic["herb_match_warning"] = fallback.get("herb_match_warning", "")
        diagnostic["herb_match_status"] = fallback.get("herb_match_status", "")
        diagnostic["herb_match_queries"] = fallback.get("herb_match_queries", [])
        diagnostic["matched_herb_id"] = fallback.get("matched_herb_id", "")
        diagnostic["matched_herb_name"] = fallback.get("matched_herb_name", "")
        diagnostic["aliases_used"] = fallback.get("aliases_used", [])
        diagnostic["raw_ingredient_count"] = len(fallback["ingredient_rows"])
        diagnostic["raw_target_relation_count"] = len(fallback["target_rows"])
        fallback_html = fallback.get("html", "")
        if fallback_html:
            selected_html_path.write_text(fallback_html, encoding="utf-8")
            write_herb_persistent_cache(
                project_root=project_root,
                herb_name=herb_name,
                detail_url=detail_url,
                html=fallback_html,
                ingredient_rows=fallback["ingredient_rows"],
                target_rows=fallback["target_rows"],
                source=fallback["source"],
            )
        write_tcmsp_debug_meta(out_root, diagnostic)
        return fallback["ingredient_rows"], fallback["target_rows"], detail_url, fallback_html, diagnostic

    diagnostic["fallback_searched_paths"] = searched_paths
    diagnostic["final_decision"] = "failed_no_valid_remote_or_cache"
    diagnostic["tcmsp_data_source"] = "unavailable"
    diagnostic["tcmsp_source_conclusion"] = "TCMSP unavailable; supplemental HerbiV/SymMap2 sources should be attempted."
    diagnostic["failure_reason"] = diagnostic.get("failure_reason") or "no_valid_remote_or_cache"
    diagnostic["source_completeness_level"] = "tcmsp_unavailable"
    diagnostic["raw_ingredient_count"] = 0
    diagnostic["raw_target_relation_count"] = 0
    write_tcmsp_debug_meta(out_root, diagnostic)
    write_json(out_root / "01_ingredients" / "tcmsp_fetch_diagnostic.json", diagnostic)
    return [], [], detail_url, "", diagnostic


def find_molecule_record(compound_name: str) -> dict:
    response = http_get_tcmsp(
        f"{TCMSP_BASE}/tcmspsearch.php",
        params={
            "qs": "molecule_name",
            "q": compound_name,
            "token": "fae1b3f834b798901e77ee835d7e468f",
        },
        timeout=40,
    )
    if response.status_code != 200:
        fail(f"TCMSP molecule search returned HTTP {response.status_code}.")
    data_index = response.text.find("data:")
    if data_index == -1:
        fail(f"TCMSP molecule search did not return a data block for {compound_name}.")
    records = extract_balanced_array(response.text, data_index)
    if not records:
        fail(f"TCMSP molecule search did not find compound: {compound_name}")
    wanted = compound_name.strip().lower()
    exact = [
        row
        for row in records
        if str(row.get("molecule_name", "")).strip().lower() == wanted
        or wanted
        in {
            part.strip().lower()
            for part in (
                str(row.get("synonyms", "")) + "," + str(row.get("molecule_synonyms", ""))
            ).split(",")
        }
    ]
    return exact[0] if exact else records[0]


def extract_js_var_array(html: str, var_name: str) -> list[dict]:
    marker = f"var {var_name}"
    marker_index = html.find(marker)
    if marker_index == -1:
        return []
    data = extract_balanced_array(html, marker_index)
    return data if isinstance(data, list) else []


def normalize_compound_input(name: str) -> str:
    return COMPOUND_NAME_ALIASES.get(name.strip(), name.strip())


PUBCHEM_SYNONYMS = {
    "dehydrotanshinone ii a": ["dehydrotanshinone iia"],
    "tanshinone ii a": ["tanshinone iia"],
    "tanshinone ⅱa": ["tanshinone iia"],
    "cryptotanshinone": ["cryptotanshinone"],
    "miltirone": ["miltirone"],
    "luteolin": ["luteolin"],
    "salvianolic acid b": ["salvianolic acid B"],
    "salvianolic acid a": ["salvianolic acid A"],
    "salvianolic acid g": ["salvianolic acid G"],
    "dihydrotanshinone i": ["dihydrotanshinone I"],
    "1,2,5,6-tetrahydrotanshinone": ["dihydrotanshinone I", "1,2,5,6-tetrahydrotanshinone"],
    "isocryptotanshi-none": ["isocryptotanshinone"],
    "baicalin": ["baicalin"],
    "isoimperatorin": ["isoimperatorin"],
    "sugiol": ["sugiol"],
    "kushenol j": ["Kushenol J", "Kushenol-J", "KushenolJ"],
    "(-)-maackiain-3-o-glucosyl-6'-o-malonate": [
        "Maackiain-3-O-glucosyl-6'-O-malonate",
        "maackiain glucoside",
        "maackiain-3-O-glucoside",
        "maackiain malonylglucoside",
        "maackiain-3-O-glucosyl-6-O-malonate",
    ],
    "maackiain-3-o-glucosyl-6'-o-malonate": [
        "maackiain glucoside",
        "maackiain-3-O-glucoside",
        "maackiain malonylglucoside",
        "maackiain-3-O-glucosyl-6-O-malonate",
    ],
}


MOJIBAKE_REPLACEMENTS = {
    "Î‘": "α",
    "Î’": "β",
    "Î±": "α",
    "Î²": "β",
    "Î³": "γ",
    "â†’": "→",
    "â€“": "-",
    "â€”": "-",
    "âˆ’": "-",
    "â€™": "'",
    "â€˜": "'",
    "â€œ": '"',
    "â€": '"',
    "Â ": " ",
    "Â": "",
}


def repair_mojibake_text(text: str) -> tuple[str, bool]:
    value = str(text or "")
    repaired = value
    for old, new in MOJIBAKE_REPLACEMENTS.items():
        repaired = repaired.replace(old, new)
    if "Ã" in repaired or "Â" in repaired:
        try:
            candidate = repaired.encode("latin1", errors="ignore").decode("utf-8", errors="ignore")
            if candidate and len(candidate) >= len(repaired) * 0.5:
                repaired = candidate
        except Exception:
            pass
    repaired = re.sub(r"\s+", " ", repaired).strip()
    return repaired, repaired != value


def normalize_ingredient_display_name(name: str) -> dict[str, str]:
    repaired, changed = repair_mojibake_text(name)
    normalized = normalize_pubchem_name(repaired)
    return {
        "raw_ingredient_name": str(name or "").strip(),
        "ingredient_name": repaired or str(name or "").strip(),
        "normalized_ingredient_name": normalized,
        "name_repair_applied": str(bool(changed)).lower(),
    }


def normalize_pubchem_name(name: str) -> str:
    value, _ = repair_mojibake_text(name)
    value = value.strip()
    replacements = {
        "Ⅰ": "I",
        "Ⅱ": "II",
        "Ⅲ": "III",
        "Ⅳ": "IV",
        "Ⅴ": "V",
        "α": "alpha",
        "β": "beta",
        "γ": "gamma",
        "→": "->",
        "–": "-",
        "—": "-",
        "−": "-",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = re.sub(r"_qt$", "", value, flags=re.IGNORECASE)
    value = re.sub(
        r"\b([A-Za-z][A-Za-z0-9-]*),\s*([A-Za-z])\b",
        lambda match: f"{match.group(1)} {match.group(2).upper()}",
        value,
    )
    value = re.sub(r"\s+", " ", value)
    value = value.replace("isocryptotanshi-none", "isocryptotanshinone")
    value = re.sub(r"\bII\s+A\b", "IIA", value, flags=re.IGNORECASE)
    return value.strip()


def split_pubchem_aliases(name: str) -> list[str]:
    parts = re.split(r"\s*(?:;|/|\bor\b|\band\b)\s*", name, flags=re.IGNORECASE)
    output = []
    for part in parts:
        part = part.strip(" ,")
        if part and len(part) >= 3:
            output.append(part)
    return output


def skipped_pubchem_candidates(name: str) -> list[tuple[str, str]]:
    cleaned = normalize_pubchem_name(name)
    skipped = []
    aglycone = re.sub(
        r"[- ]?(?:O-)?(?:alpha|beta|α|β)?-?[DL]?-?(?:gluco|galacto|rhamno|xylopyrano|glucopyrano)[A-Za-z0-9'(),→\\ ]*-?",
        "",
        cleaned,
        flags=re.IGNORECASE,
    ).strip(" -")
    if aglycone and aglycone != cleaned and len(aglycone) >= 4:
        skipped.append((aglycone, "skipped_aglycone_candidate_not_equivalent_to_original"))
    return skipped


def pubchem_query_candidates(name: str) -> list[tuple[str, str]]:
    cleaned = normalize_pubchem_name(name)
    without_stereo = re.sub(r"^\(\s*[+-]\s*\)-", "", cleaned).strip()
    candidates: list[tuple[str, str]] = [(name.strip(), "raw_name"), (cleaned, "normalized_name"), (cleaned.lower(), "normalized_lower")]
    for alias in split_pubchem_aliases(cleaned):
        candidates.append((alias, "split_alias"))
    if without_stereo and without_stereo != cleaned:
        candidates.extend([(without_stereo, "without_stereochemical_prefix"), (without_stereo.lower(), "without_stereochemical_prefix_lower")])
    candidates.extend((item, "curated_synonym") for item in PUBCHEM_SYNONYMS.get(cleaned.lower(), []))
    candidates.extend((item, "curated_synonym") for item in PUBCHEM_SYNONYMS.get(without_stereo.lower(), []))
    output = []
    seen = set()
    for candidate, kind in candidates:
        candidate = candidate.strip()
        key = candidate.lower()
        if not candidate or key in seen:
            continue
        seen.add(key)
        output.append((candidate, kind))
    return output


def empty_pubchem_result(status: str = "not_queried", query_name: str = "", error: str = "") -> dict[str, str]:
    return {
        "pubchem_cid": "",
        "smiles": "",
        "canonical_smiles": "",
        "isomeric_smiles": "",
        "inchikey": "",
        "inchi_key": "",
        "pubchem_match_status": status,
        "pubchem_query_name": query_name,
        "pubchem_attempted_names": query_name,
        "pubchem_error": error,
    }


def fetch_pubchem_structure_by_cid(
    cid: str,
    cache: dict[str, dict[str, str]],
    attempts: list[dict] | None = None,
    *,
    ingredient_id: str = "",
    ingredient_name: str = "",
    raw_ingredient_name: str = "",
    cid_origin_source: str = "",
) -> dict[str, str]:
    cid = str(cid or "").strip()
    if not cid:
        return empty_pubchem_result("not_queried", "", "empty_cid")
    key = cache_key(f"cid:{cid}")
    if key in cache:
        result = dict(cache[key])
        if not pubchem_cached_result_is_retryable(result):
            if attempts is not None:
                attempts.append(
                    {
                        "ingredient_id": ingredient_id,
                        "raw_ingredient_name": raw_ingredient_name or ingredient_name,
                        "ingredient_name": ingredient_name,
                        "candidate": cid,
                        "candidate_type": f"source_cid:{cid_origin_source}",
                        "query_url": "",
                        "http_status": "",
                        "accepted": "true" if result.get("smiles") or result.get("inchikey") else "false",
                        "decision_reason": "cache_hit_by_cid",
                        "pubchem_cid": cid,
                        "inchikey": result.get("inchikey", ""),
                    }
                )
            return result
        if attempts is not None:
            attempts.append(
                {
                    "ingredient_id": ingredient_id,
                    "raw_ingredient_name": raw_ingredient_name or ingredient_name,
                    "ingredient_name": ingredient_name,
                    "candidate": cid,
                    "candidate_type": f"source_cid:{cid_origin_source}",
                    "query_url": "",
                    "http_status": "",
                    "accepted": "false",
                    "decision_reason": "retrying_transient_cache_miss_by_cid",
                    "pubchem_cid": cid,
                    "inchikey": result.get("inchikey", ""),
                }
            )
    url = (
        "https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/cid/"
        f"{quote(cid)}/property/CanonicalSMILES,IsomericSMILES,InChIKey/JSON"
    )
    attempt_row = {
        "ingredient_id": ingredient_id,
        "raw_ingredient_name": raw_ingredient_name or ingredient_name,
        "ingredient_name": ingredient_name,
        "candidate": cid,
        "candidate_type": f"source_cid:{cid_origin_source}",
        "query_url": url,
        "http_status": "",
        "accepted": "false",
        "decision_reason": "",
        "pubchem_cid": cid,
        "inchikey": "",
    }
    try:
        response = http_get(url, timeout=8, retries=1)
        attempt_row["http_status"] = response.status_code
        if response.status_code != 200:
            attempt_row["decision_reason"] = f"http_{response.status_code}"
            if attempts is not None:
                attempts.append(attempt_row)
            result = empty_pubchem_result("source_cid_property_lookup_failed", cid, attempt_row["decision_reason"])
            result["pubchem_cid"] = cid
            if not pubchem_error_is_transient(attempt_row["decision_reason"]):
                cache[key] = result
            return result
        props = response.json().get("PropertyTable", {}).get("Properties", [])
        if not props:
            attempt_row["decision_reason"] = "no_properties"
            if attempts is not None:
                attempts.append(attempt_row)
            result = empty_pubchem_result("source_cid_no_properties", cid, "no_properties")
            result["pubchem_cid"] = cid
            cache[key] = result
            return result
        first = props[0]
        canonical = first.get("CanonicalSMILES") or first.get("SMILES") or first.get("ConnectivitySMILES") or ""
        isomeric = first.get("IsomericSMILES") or first.get("SMILES") or canonical
        inchikey = first.get("InChIKey", "")
        result = {
            "pubchem_cid": str(first.get("CID", cid)),
            "smiles": canonical or isomeric,
            "canonical_smiles": canonical,
            "isomeric_smiles": isomeric,
            "inchikey": inchikey,
            "inchi_key": inchikey,
            "pubchem_match_status": "matched_source_pubchem_cid",
            "pubchem_query_name": cid,
            "pubchem_attempted_names": cid,
            "pubchem_error": "",
        }
        attempt_row["accepted"] = "true"
        attempt_row["decision_reason"] = "accepted_by_source_cid"
        attempt_row["inchikey"] = inchikey
        if attempts is not None:
            attempts.append(attempt_row)
        cache[key] = result
        return result
    except Exception as exc:  # noqa: BLE001
        attempt_row["decision_reason"] = f"{type(exc).__name__}: {exc}"
        if attempts is not None:
            attempts.append(attempt_row)
        result = empty_pubchem_result("source_cid_property_lookup_failed", cid, attempt_row["decision_reason"])
        result["pubchem_cid"] = cid
        if not pubchem_error_is_transient(attempt_row["decision_reason"]):
            cache[key] = result
        return result


def load_pubchem_cache(cache_path: Path) -> dict[str, dict[str, str]]:
    if not cache_path.exists():
        return {}
    data = read_json(cache_path)
    return data if isinstance(data, dict) else {}


def pubchem_result_has_structure(result: dict[str, str] | None) -> bool:
    if not isinstance(result, dict):
        return False
    return bool(
        str(result.get("smiles", "")).strip()
        or str(result.get("canonical_smiles", "")).strip()
        or str(result.get("isomeric_smiles", "")).strip()
        or str(result.get("inchikey", "")).strip()
        or str(result.get("inchi_key", "")).strip()
    )


def pubchem_error_is_transient(error: str) -> bool:
    error = str(error or "").lower()
    transient_markers = (
        "http 5",
        "http_5",
        "503",
        "502",
        "504",
        "timeout",
        "timed out",
        "connection",
        "temporarily",
        "remote end closed",
        "max retries",
    )
    return any(marker in error for marker in transient_markers)


def pubchem_cached_result_is_retryable(result: dict[str, str] | None) -> bool:
    if not isinstance(result, dict):
        return False
    if pubchem_result_has_structure(result):
        return False
    status = str(result.get("pubchem_match_status", "")).lower()
    error = str(result.get("pubchem_error", ""))
    return status in {"source_cid_property_lookup_failed", "remote_unavailable"} or pubchem_error_is_transient(error)


def merge_pubchem_cache(
    primary: dict[str, dict[str, str]],
    fallback: dict[str, dict[str, str]],
) -> dict[str, dict[str, str]]:
    for key, value in fallback.items():
        if not isinstance(value, dict):
            continue
        existing = primary.get(key)
        if existing is None or (pubchem_result_has_structure(value) and not pubchem_result_has_structure(existing)):
            primary[key] = value
    return primary


def update_shared_pubchem_cache(shared_cache_path: Path | None, pubchem_cache: dict[str, dict[str, str]]) -> None:
    if shared_cache_path is None:
        return
    shared_cache = load_pubchem_cache(shared_cache_path)
    for key, value in pubchem_cache.items():
        if not isinstance(value, dict) or not pubchem_result_has_structure(value):
            continue
        existing = shared_cache.get(key)
        if existing is None or not pubchem_result_has_structure(existing):
            shared_cache[key] = value
    write_json(shared_cache_path, shared_cache)


def cache_key(name: str) -> str:
    return hashlib.sha1(name.encode("utf-8")).hexdigest()


def fetch_pubchem_structure(
    name: str,
    cache: dict[str, dict[str, str]],
    attempts: list[dict] | None = None,
    *,
    ingredient_id: str = "",
    raw_ingredient_name: str = "",
) -> dict[str, str]:
    key = cache_key(name)
    if key in cache:
        cached = dict(cache[key])
        if not pubchem_cached_result_is_retryable(cached):
            if attempts is not None:
                attempts.append(
                    {
                        "ingredient_id": ingredient_id,
                        "raw_ingredient_name": raw_ingredient_name or name,
                        "ingredient_name": name,
                        "candidate": cached.get("pubchem_query_name", ""),
                        "candidate_type": "cache_hit",
                        "query_url": "",
                        "http_status": "",
                        "accepted": "true" if cached.get("pubchem_cid") or cached.get("smiles") else "false",
                        "decision_reason": "cache_hit",
                        "pubchem_cid": cached.get("pubchem_cid", ""),
                        "inchikey": cached.get("inchikey", ""),
                    }
                )
            return cached
        if attempts is not None:
            attempts.append(
                {
                    "ingredient_id": ingredient_id,
                    "raw_ingredient_name": raw_ingredient_name or name,
                    "ingredient_name": name,
                    "candidate": cached.get("pubchem_query_name", ""),
                    "candidate_type": "cache_hit",
                    "query_url": "",
                    "http_status": "",
                    "accepted": "false",
                    "decision_reason": "retrying_transient_cache_miss",
                    "pubchem_cid": cached.get("pubchem_cid", ""),
                    "inchikey": cached.get("inchikey", ""),
                }
            )

    errors: list[str] = []
    candidates = pubchem_query_candidates(name)
    attempted = "; ".join(candidate for candidate, _kind in candidates)
    for candidate, candidate_type in candidates:
        url = (
            "https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/name/"
            f"{quote(candidate)}/property/CanonicalSMILES,IsomericSMILES,InChIKey/JSON"
        )
        attempt_row = {
            "ingredient_id": ingredient_id,
            "raw_ingredient_name": raw_ingredient_name or name,
            "ingredient_name": name,
            "candidate": candidate,
            "candidate_type": candidate_type,
            "query_url": url,
            "http_status": "",
            "accepted": "false",
            "decision_reason": "",
            "pubchem_cid": "",
            "inchikey": "",
        }
        try:
            response = http_get(url, timeout=8, retries=1)
            attempt_row["http_status"] = response.status_code
            if response.status_code == 404:
                errors.append(f"{candidate}: not_found")
                attempt_row["decision_reason"] = "not_found"
                if attempts is not None:
                    attempts.append(attempt_row)
                continue
            if response.status_code != 200:
                errors.append(f"{candidate}: HTTP {response.status_code}")
                attempt_row["decision_reason"] = f"http_{response.status_code}"
                if attempts is not None:
                    attempts.append(attempt_row)
                continue
            data = response.json()
            props = data.get("PropertyTable", {}).get("Properties", [])
            if not props:
                errors.append(f"{candidate}: no_properties")
                attempt_row["decision_reason"] = "no_properties"
                if attempts is not None:
                    attempts.append(attempt_row)
                continue
            first = props[0]
            canonical = (
                first.get("CanonicalSMILES")
                or first.get("SMILES")
                or first.get("ConnectivitySMILES")
                or ""
            )
            isomeric = first.get("IsomericSMILES") or first.get("SMILES") or canonical
            inchikey = first.get("InChIKey", "")
            result = {
                "pubchem_cid": str(first.get("CID", "")),
                "smiles": canonical or isomeric,
                "canonical_smiles": canonical,
                "isomeric_smiles": isomeric,
                "inchikey": inchikey,
                "inchi_key": inchikey,
                "pubchem_match_status": "matched_exact" if candidate == name.strip() else "matched_fallback",
                "pubchem_query_name": candidate,
                "pubchem_attempted_names": attempted,
                "pubchem_error": "",
            }
            attempt_row["accepted"] = "true"
            attempt_row["decision_reason"] = "accepted"
            attempt_row["pubchem_cid"] = result["pubchem_cid"]
            attempt_row["inchikey"] = result["inchikey"]
            if attempts is not None:
                attempts.append(attempt_row)
                for skipped_candidate, reason in skipped_pubchem_candidates(name):
                    attempts.append(
                        {
                            "ingredient_id": ingredient_id,
                            "raw_ingredient_name": raw_ingredient_name or name,
                            "ingredient_name": name,
                            "candidate": skipped_candidate,
                            "candidate_type": "skipped_structural_variant",
                            "query_url": "",
                            "http_status": "",
                            "accepted": "false",
                            "decision_reason": reason,
                            "pubchem_cid": "",
                            "inchikey": "",
                        }
                    )
            cache[key] = result
            return result
        except Exception as exc:
            errors.append(f"{candidate}: {type(exc).__name__}: {exc}")
            attempt_row["decision_reason"] = f"{type(exc).__name__}: {exc}"
            if attempts is not None:
                attempts.append(attempt_row)

    transient_failure = any(pubchem_error_is_transient(error) for error in errors)
    result = empty_pubchem_result("remote_unavailable" if transient_failure else "not_found", "", "; ".join(errors[:5]))
    result["pubchem_attempted_names"] = attempted
    if attempts is not None:
        for skipped_candidate, reason in skipped_pubchem_candidates(name):
            attempts.append(
                {
                    "ingredient_id": ingredient_id,
                    "raw_ingredient_name": raw_ingredient_name or name,
                    "ingredient_name": name,
                    "candidate": skipped_candidate,
                    "candidate_type": "skipped_structural_variant",
                    "query_url": "",
                    "http_status": "",
                    "accepted": "false",
                    "decision_reason": reason,
                    "pubchem_cid": "",
                    "inchikey": "",
                }
            )
    if not transient_failure:
        cache[key] = result
    return result


def normalize_target_query(target_name: str) -> str:
    return re.sub(r"\s+", " ", target_name.strip())


def make_mapping_record(
    target_name: str,
    *,
    gene: str = "",
    uniprot: str = "",
    mapping_status: str = "unmapped",
    normalized_query: str | None = None,
    attempted_queries: list[str] | None = None,
    failure_reason: str = "",
    suggested_gene_symbol: str = "",
    mapping_note: str = "",
) -> dict[str, str]:
    return {
        "target_name": target_name,
        "target_gene": gene,
        "uniprot_id": uniprot,
        "mapping_status": mapping_status,
        "normalized_query": normalized_query if normalized_query is not None else normalize_target_query(target_name),
        "attempted_queries": "; ".join(attempted_queries or []),
        "failure_reason": failure_reason,
        "suggested_gene_symbol": suggested_gene_symbol,
        "mapping_note": mapping_note,
    }


def map_target_record(target_name: str, cache: dict[str, dict[str, str]]) -> dict[str, str]:
    normalized_query = normalize_target_query(target_name)
    if target_name in TARGET_NAME_TO_GENE:
        gene, uniprot = TARGET_NAME_TO_GENE[target_name]
        return make_mapping_record(
            target_name,
            gene=gene,
            uniprot=uniprot,
            mapping_status="mapped_alias",
            normalized_query=normalized_query,
            attempted_queries=[target_name],
        )
    if target_name in CANONICAL_AMBIGUOUS_TARGETS:
        item = CANONICAL_AMBIGUOUS_TARGETS[target_name]
        return make_mapping_record(
            target_name,
            gene=item["gene"],
            uniprot=item["uniprot"],
            mapping_status="mapped_canonical_from_ambiguous",
            normalized_query=normalized_query,
            attempted_queries=[target_name],
            suggested_gene_symbol=item["suggested_gene_symbol"],
            mapping_note=item["note"],
        )
    if target_name in AMBIGUOUS_TARGETS:
        item = AMBIGUOUS_TARGETS[target_name]
        return make_mapping_record(
            target_name,
            mapping_status="ambiguous",
            normalized_query=normalized_query,
            attempted_queries=[target_name],
            failure_reason="ambiguous_complex_or_multi_gene_target",
            suggested_gene_symbol=item["suggested_gene_symbol"],
            mapping_note=item["note"],
        )
    if target_name in IGNORED_TARGET_NAMES:
        return make_mapping_record(
            target_name,
            mapping_status="ignored_non_human_or_irrelevant",
            normalized_query=normalized_query,
            attempted_queries=[target_name],
            failure_reason="non_human_or_not_actionable_for_human_target_network",
        )
    if target_name in TARGET_NAME_TO_GENE:
        gene, uniprot = TARGET_NAME_TO_GENE[target_name]
        return make_mapping_record(target_name, gene=gene, uniprot=uniprot, mapping_status="mapped_alias")
    if target_name in cache:
        return cache[target_name]

    query = f'({target_name}) AND (organism_id:9606)'
    attempted_queries = [query]
    try:
        response = http_get(
            "https://rest.uniprot.org/uniprotkb/search",
            params={
                "query": query,
                "format": "json",
                "fields": "accession,gene_primary,protein_name,reviewed",
                "size": 1,
            },
            timeout=25,
            retries=1,
        )
        if response.status_code != 200:
            cache[target_name] = make_mapping_record(
                target_name,
                mapping_status="unmapped",
                normalized_query=normalized_query,
                attempted_queries=attempted_queries,
                failure_reason=f"uniprot_http_{response.status_code}",
            )
            return cache[target_name]
        results = response.json().get("results", [])
        if not results:
            cache[target_name] = make_mapping_record(
                target_name,
                mapping_status="unmapped",
                normalized_query=normalized_query,
                attempted_queries=attempted_queries,
                failure_reason="no_uniprot_result",
            )
            return cache[target_name]
        entry = results[0]
        genes = entry.get("genes") or []
        gene_symbol = ""
        if genes:
            gene_symbol = genes[0].get("geneName", {}).get("value", "")
        cache[target_name] = make_mapping_record(
            target_name,
            gene=gene_symbol,
            uniprot=entry.get("primaryAccession", ""),
            mapping_status="mapped_uniprot" if gene_symbol else "unmapped",
            normalized_query=normalized_query,
            attempted_queries=attempted_queries,
            failure_reason="" if gene_symbol else "uniprot_result_without_gene_symbol",
        )
        time.sleep(0.05)
        return cache[target_name]
    except Exception:
        cache[target_name] = make_mapping_record(
            target_name,
            mapping_status="unmapped",
            normalized_query=normalized_query,
            attempted_queries=attempted_queries,
            failure_reason="uniprot_request_failed",
        )
        return cache[target_name]


def map_target_name(target_name: str, cache: dict[str, dict[str, str]]) -> tuple[str, str]:
    record = map_target_record(target_name, cache)
    return record.get("target_gene", ""), record.get("uniprot_id", "")


def parse_optional_float(value: object) -> float | None:
    text = str(value if value is not None else "").strip()
    if not text or text.lower() in {"none", "nan", "null", "na", "n/a"}:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _strict_true(value: object) -> bool:
    return value is True or (
        isinstance(value, str) and value.strip().casefold() in {"true", "1", "yes"}
    )


def classify_target_evidence(value: object) -> str:
    """Classify, but never upgrade, the source-provided evidence label."""
    text = str(value or "").strip().casefold()
    if not text:
        return "source_record_unspecified"
    if any(token in text for token in ("predict", "computational", "svm", "random forest")):
        return "computational_prediction"
    if any(
        token in text
        for token in ("experimental", "experimentally", "assay", "in vitro", "in_vitro")
    ):
        return "experimental_source_label"
    if any(token in text for token in ("literature", "curated")):
        return "literature_curated_source_label"
    if "related_component" in text or "related component" in text:
        return "database_related_component"
    return "source_record_unclassified"


def target_edge_eligibility(row: dict) -> tuple[bool, str]:
    """Apply edge-level eligibility; ingredient eligibility cannot override it."""
    status = str(
        row.get("network_eligibility")
        or row.get("source_eligibility")
        or row.get("herbiv_category")
        or ""
    ).strip().casefold()
    reason = str(
        row.get("network_eligibility_reason")
        or row.get("source_eligibility_reason")
        or ""
    ).strip()
    if status in {
        "true",
        "eligible",
        "candidate",
        "selected",
        "main_network",
        "pending_activity_filter",
    }:
        return True, reason or f"eligible_edge_status_{status}"
    if status in {"exploratory", "excluded", "false", "ineligible"}:
        return False, reason or f"excluded_edge_status_{status}"

    sources = {
        part.strip()
        for part in str(row.get("eligible_source_list") or row.get("source_list") or row.get("source") or "").split(";")
        if part.strip()
    }
    if "SymMap2" in sources:
        return True, "eligible_symmap_direct_target_relation"
    if "TCMSP" in sources:
        return True, "eligible_after_tcmsp_ingredient_filter"
    if "HerbiV" in sources:
        category = str(row.get("herbiv_category") or row.get("source_eligibility") or "").strip().casefold()
        if category == "candidate":
            return True, "eligible_herbiv_candidate_score"
        return False, "excluded_herbiv_not_candidate"
    return False, "excluded_unknown_edge_source_policy"


def _source_set(value: object) -> set[str]:
    return {
        part.strip()
        for part in str(value or "").split(";")
        if part.strip()
    }


def ingredient_source_eligibility(
    row: dict, filters: dict
) -> dict[str, tuple[bool, str]]:
    """Return source-local ingredient decisions before their OR aggregation."""
    sources = {
        part.strip()
        for part in str(row.get("source_list", "") or row.get("source", "")).split(";")
        if part.strip()
    }
    ob = parse_optional_float(row.get("ob"))
    dl = parse_optional_float(row.get("dl"))
    ob_min = float(filters.get("ob_min", 30))
    dl_min = float(filters.get("dl_min", 0.18))
    decisions: dict[str, tuple[bool, str]] = {}
    if "TCMSP" in sources:
        if ob is not None and dl is not None and ob >= ob_min and dl >= dl_min:
            decisions["TCMSP"] = (True, "eligible_tcmsp_ob_dl")
        else:
            decisions["TCMSP"] = (False, "excluded_tcmsp_ob_dl_threshold")
    if "HerbiV" in sources:
        category = str(row.get("source_eligibility") or row.get("herbiv_category") or "").strip().lower()
        if category == "candidate":
            decisions["HerbiV"] = (True, "eligible_herbiv_candidate_score")
        elif category == "exploratory":
            decisions["HerbiV"] = (False, "excluded_herbiv_exploratory_only")
        else:
            decisions["HerbiV"] = (False, "excluded_herbiv_not_candidate")
    if "SymMap2" in sources:
        # SymMap2 contributes direct linked network evidence without inventing
        # TCMSP-style OB/DL values. Docking readiness is gated by structure later.
        decisions["SymMap2"] = (True, "eligible_symmap_direct_target_relation")
    return decisions


def ingredient_eligibility(row: dict, target_mol_ids: set[str], filters: dict) -> tuple[bool, str]:
    if row.get("MOL_ID", "") not in target_mol_ids:
        return False, "excluded_no_target_relation"
    decisions = ingredient_source_eligibility(row, filters)
    eligible_reasons = [reason for eligible, reason in decisions.values() if eligible]
    excluded_reasons = [reason for eligible, reason in decisions.values() if not eligible]
    if eligible_reasons:
        return True, "; ".join([*eligible_reasons, *excluded_reasons])
    if excluded_reasons:
        return False, "; ".join(excluded_reasons)
    return False, "excluded_unknown_source_policy"


def target_edge_eligible_sources(row: dict) -> set[str]:
    """Return only sources whose own record makes this edge eligible."""
    if "eligible_source_list" in row:
        return _source_set(row.get("eligible_source_list"))
    sources = _source_set(row.get("source_list") or row.get("source"))
    output: set[str] = set()
    for source in sources:
        if source in {"TCMSP", "SymMap2"}:
            output.add(source)
        elif source == "HerbiV" and str(
            row.get("source_eligibility") or row.get("herbiv_category") or ""
        ).strip().casefold() == "candidate":
            output.add(source)
    return output


def edge_confidence_for_sources(
    row: dict, eligible_sources: set[str]
) -> tuple[float | str, float | str]:
    """Recompute edge confidence after the TCMSP activity filter is known."""
    records = row.get("_source_edge_records")
    if not isinstance(records, list) or not records:
        return row.get("ingredient_confidence", ""), row.get("relation_confidence", "")
    ingredient_values: list[float] = []
    relation_values: list[float] = []
    defaults = {
        "TCMSP": (1.0, 0.9),
        "HerbiV": (0.75, 0.65),
        "SymMap2": (0.65, 0.60),
    }
    for record in records:
        record_sources = _source_set(record.get("source_list") or record.get("source"))
        for source in record_sources & eligible_sources:
            if source == "HerbiV" and str(
                record.get("source_eligibility") or record.get("herbiv_category") or ""
            ).strip().casefold() != "candidate":
                continue
            ingredient_default, relation_default = defaults[source]
            ingredient_value = parse_optional_float(record.get("ingredient_confidence"))
            relation_value = parse_optional_float(record.get("relation_confidence"))
            ingredient_values.append(
                ingredient_default if ingredient_value is None else ingredient_value
            )
            relation_values.append(
                relation_default if relation_value is None else relation_value
            )
    if not ingredient_values or not relation_values:
        return row.get("ingredient_confidence", ""), row.get("relation_confidence", "")
    return round(max(ingredient_values), 4), round(max(relation_values), 4)


def finalize_union_edge_audit(out_root: Path, compound_targets: list[dict]) -> None:
    """Finalize provisional edge eligibility after source-specific filtering."""
    edge_path = out_root / "02_compound_targets" / "union_ingredient_target_edges.csv"
    if not edge_path.exists():
        return
    edge_rows = read_rows(edge_path)
    if not edge_rows:
        return
    final_by_pair = {
        (
            str(row.get("ingredient_id", "") or "").strip(),
            str(row.get("target_gene", "") or "").strip().upper(),
        ): row
        for row in compound_targets
        if row.get("ingredient_id") and row.get("target_gene")
    }
    for edge in edge_rows:
        pair = (
            str(edge.get("ingredient_id", "") or "").strip(),
            str(edge.get("target_gene", "") or "").strip().upper(),
        )
        final = final_by_pair.get(pair)
        if final is not None:
            edge["eligible_source_list"] = final.get("eligible_source_list", "") or final.get(
                "source_list", ""
            )
            edge["ingredient_confidence"] = final.get("ingredient_confidence", "")
            edge["relation_confidence"] = final.get("relation_confidence", "")
            edge["network_eligibility"] = "eligible"
            edge["network_eligibility_reason"] = final.get(
                "network_eligibility_reason", "final_source_and_ingredient_filters_passed"
            )
            edge["evidence_class"] = final.get("evidence_class", edge.get("evidence_class", ""))
            edge["mapping_validated"] = final.get(
                "mapping_validated", edge.get("mapping_validated", "false")
            )
        elif str(edge.get("network_eligibility", "")).strip().casefold() not in {
            "exploratory",
            "excluded",
        }:
            edge["eligible_source_list"] = ""
            edge["network_eligibility"] = "excluded"
            previous_reason = str(edge.get("network_eligibility_reason", "") or "").strip()
            edge["network_eligibility_reason"] = "; ".join(
                part
                for part in (
                    previous_reason,
                    "excluded_by_final_source_specific_ingredient_or_edge_filter",
                )
                if part
            )
    write_rows(edge_path, edge_rows, rows_fieldnames(edge_rows))

    target_path = out_root / "02_compound_targets" / "union_targets.csv"
    if not target_path.exists():
        return
    target_rows = read_rows(target_path)
    final_confidence_by_gene: dict[str, float] = {}
    for edge in edge_rows:
        if str(edge.get("network_eligibility", "")).strip().casefold() != "eligible":
            continue
        gene = str(edge.get("target_gene", "") or "").strip().upper()
        confidence = parse_optional_float(edge.get("relation_confidence")) or 0.0
        if gene:
            final_confidence_by_gene[gene] = max(
                final_confidence_by_gene.get(gene, 0.0), confidence
            )
    for target in target_rows:
        gene = str(target.get("target_gene", "") or "").strip().upper()
        target["target_confidence"] = final_confidence_by_gene.get(gene, 0.0)
    if target_rows:
        write_rows(target_path, target_rows, rows_fieldnames(target_rows))


def is_active_ingredient(row: dict, target_mol_ids: set[str], filters: dict) -> bool:
    return ingredient_eligibility(row, target_mol_ids, filters)[0]


def collect_mapping_issues(rows: list[dict]) -> list[dict]:
    grouped: dict[tuple[str, str], dict[str, set[str] | int | str]] = {}
    for row in rows:
        target_name = row.get("target_name", "")
        if not target_name:
            continue
        key = (target_name, row.get("mapping_status", "unmapped"))
        item = grouped.setdefault(
            key,
            {
                "ingredient_ids": set(),
                "ingredient_names": set(),
                "target_row_count": 0,
                "normalized_query": row.get("normalized_query", ""),
                "attempted_queries": row.get("attempted_queries", ""),
                "failure_reason": row.get("failure_reason", ""),
                "suggested_gene_symbol": row.get("suggested_gene_symbol", ""),
                "mapping_status": row.get("mapping_status", "unmapped"),
            },
        )
        item["ingredient_ids"].add(row.get("ingredient_id", ""))  # type: ignore[union-attr]
        item["ingredient_names"].add(row.get("ingredient_name", ""))  # type: ignore[union-attr]
        item["target_row_count"] = int(item["target_row_count"]) + 1
    output = []
    for (target_name, _status), item in sorted(grouped.items()):
        output.append(
            {
                "target_name": target_name,
                "normalized_query": item["normalized_query"],
                "attempted_queries": item["attempted_queries"],
                "failure_reason": item["failure_reason"],
                "suggested_gene_symbol": item["suggested_gene_symbol"],
                "mapping_status": item["mapping_status"],
                "ingredient_ids": "; ".join(sorted(value for value in item["ingredient_ids"] if value)),  # type: ignore[index]
                "ingredient_names": "; ".join(sorted(value for value in item["ingredient_names"] if value)),  # type: ignore[index]
                "target_row_count": item["target_row_count"],
                "source": "TCMSP",
            }
        )
    return output


def collect_unmapped_targets(unmapped_rows: list[dict]) -> list[dict]:
    return collect_mapping_issues(unmapped_rows)


def build_ligand_ready_rows(active_rows: list[dict], compound_targets: list[dict]) -> list[dict]:
    target_count_by_mol: dict[str, int] = {}
    for row in compound_targets:
        target_count_by_mol[row["ingredient_id"]] = target_count_by_mol.get(row["ingredient_id"], 0) + 1
    ligand_ready_rows = []
    for row in active_rows:
        target_count = target_count_by_mol.get(row["ingredient_id"], 0)
        has_structure = bool(str(row.get("smiles", "")).strip() or str(row.get("canonical_smiles", "")).strip() or str(row.get("isomeric_smiles", "")).strip())
        if has_structure and target_count > 0:
            ligand_ready_rows.append(
                {
                    **row,
                    "compound_target_count": target_count,
                    "structure_source": row.get("smiles_origin") or "PubChem",
                    "recommended_for_docking": "true" if row.get("structure_status") in {"smiles_ready", "inchikey_and_smiles_ready"} else "false",
                }
            )
    return ligand_ready_rows


def structure_status_for_row(structure: dict[str, str], *, source_pubchem_cid: str = "") -> tuple[str, str, str, str]:
    has_smiles = bool(
        str(structure.get("smiles", "")).strip()
        or str(structure.get("canonical_smiles", "")).strip()
        or str(structure.get("isomeric_smiles", "")).strip()
    )
    has_inchikey = bool(str(structure.get("inchikey", "")).strip() or str(structure.get("inchi_key", "")).strip())
    has_cid = bool(str(structure.get("pubchem_cid", "")).strip() or source_pubchem_cid)
    if has_smiles and has_inchikey:
        return (
            "inchikey_and_smiles_ready",
            "SMILES and InChIKey available from PubChem CID property lookup; no PDBQT generated in this workflow.",
            "PubChemCID",
            "PubChemCID",
        )
    if has_smiles:
        return (
            "smiles_ready",
            "SMILES available from PubChem CID property lookup; no PDBQT generated in this workflow.",
            "PubChemCID",
            "",
        )
    if has_cid:
        return (
            "pubchem_cid_ready",
            "Source database supplied PubChem CID, but SMILES/InChIKey were not retrieved; not docking-ready.",
            "",
            "",
        )
    return ("no_structure", "No PubChem CID, SMILES, or structure file available; not docking-ready.", "", "")


def ligand_identity(row: dict) -> str:
    for field in ("inchikey", "pubchem_cid", "canonical_smiles", "isomeric_smiles"):
        value = str(row.get(field, "")).strip()
        if value:
            return f"{field}:{value}"
    fallback = str(row.get("ingredient_id", "")).strip() or str(row.get("ingredient_name", "")).strip()
    return f"ingredient:{fallback}"


def ligand_id_from_row(row: dict) -> str:
    if row.get("inchikey"):
        return row["inchikey"]
    if row.get("pubchem_cid"):
        return f"CID:{row['pubchem_cid']}"
    smiles = row.get("canonical_smiles") or row.get("isomeric_smiles") or row.get("smiles", "")
    if smiles:
        return f"SMILES:{hashlib.sha1(smiles.encode('utf-8')).hexdigest()[:12]}"
    return f"LIGAND:{row.get('ingredient_id', '')}"


def normalized_ingredient_identity(name: str) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", str(name or "").strip().lower())


def classify_multi_name(raw_name: str) -> tuple[str, list[str], str]:
    parts = [part.strip() for part in re.split(r"[,，;；/|]", str(raw_name or "")) if part.strip()]
    if len(parts) <= 1:
        return "", parts, ""
    normalized_parts = [normalized_ingredient_identity(part.replace("β", "beta").replace("Î’", "beta").replace("Î²", "beta")) for part in parts]
    first = normalized_parts[0]
    if all(first in part or part in first for part in normalized_parts[1:]):
        return "alias_multi_name", parts, "parts appear to be aliases of the same compound"
    return "ambiguous_multi_name", parts, "multiple distinct compound names in one database field"


def active_ingredient_identity(row: dict) -> tuple[str, str]:
    for field, label in (("pubchem_cid", "pubchem_cid"), ("inchikey", "inchikey")):
        value = str(row.get(field, "")).strip()
        if value:
            return label, value
    name_key = normalized_ingredient_identity(row.get("normalized_ingredient_name", "") or row.get("ingredient_name", ""))
    if name_key:
        return "normalized_ingredient_name", name_key
    fallback = str(row.get("ingredient_id", "")).strip()
    return "ingredient_id", fallback


def collapse_deduplicated_active_rows(active_rows: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for row in active_rows:
        identity_type, identity_value = active_ingredient_identity(row)
        key = f"{identity_type}:{identity_value}"
        item = grouped.setdefault(
            key,
            {
                "identity_key": identity_value,
                "identity_type": identity_type,
                "representative": row,
                "source_ingredient_ids": set(),
                "source_ingredient_names": set(),
                "source_herb_names": set(),
                "pubchem_match_status": set(),
            },
        )
        item["source_ingredient_ids"].add(row.get("ingredient_id", ""))
        item["source_ingredient_names"].add(row.get("ingredient_name", ""))
        item["source_herb_names"].add(row.get("herb_name", ""))
        item["pubchem_match_status"].add(row.get("pubchem_match_status", ""))
    output = []
    for index, item in enumerate(grouped.values(), start=1):
        row = item["representative"]
        output.append(
            {
                "deduplicated_ingredient_id": f"DAI{index:04d}",
                "identity_key": item["identity_key"],
                "identity_type": item["identity_type"],
                "ingredient_id": row.get("ingredient_id", ""),
                "raw_ingredient_name": row.get("raw_ingredient_name", ""),
                "ingredient_name": row.get("ingredient_name", ""),
                "normalized_ingredient_name": row.get("normalized_ingredient_name", ""),
                "name_repair_applied": row.get("name_repair_applied", ""),
                "pubchem_cid": row.get("pubchem_cid", ""),
                "smiles": row.get("smiles", ""),
                "canonical_smiles": row.get("canonical_smiles", ""),
                "isomeric_smiles": row.get("isomeric_smiles", ""),
                "inchikey": row.get("inchikey", ""),
                "source_ingredient_ids": "; ".join(sorted(value for value in item["source_ingredient_ids"] if value)),
                "source_ingredient_names": "; ".join(sorted(value for value in item["source_ingredient_names"] if value)),
                "source_herb_names": "; ".join(sorted(value for value in item["source_herb_names"] if value)),
                "active_record_count": len([value for value in item["source_ingredient_ids"] if value]),
                "pubchem_match_status": "; ".join(sorted(value for value in item["pubchem_match_status"] if value)),
                "ob": row.get("ob", ""),
                "dl": row.get("dl", ""),
                "source": row.get("source", ""),
                "tcmsp_data_source": row.get("tcmsp_data_source", ""),
            }
        )
    output.sort(key=lambda item: (item["identity_type"], item["identity_key"]))
    return output


def collapse_unique_ligand_ready_rows(ligand_ready_rows: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for row in ligand_ready_rows:
        key = ligand_identity(row)
        item = grouped.setdefault(
            key,
            {
                "ligand_id": ligand_id_from_row(row),
                "pubchem_cid": row.get("pubchem_cid", ""),
                "inchikey": row.get("inchikey", ""),
                "canonical_smiles": row.get("canonical_smiles", "") or row.get("smiles", ""),
                "isomeric_smiles": row.get("isomeric_smiles", ""),
                "source_ingredient_ids": set(),
                "source_ingredient_names": set(),
                "source_herb_names": set(),
                "compound_target_count": 0,
                "pubchem_match_status": set(),
            },
        )
        item["source_ingredient_ids"].add(row.get("ingredient_id", ""))
        item["source_ingredient_names"].add(row.get("ingredient_name", ""))
        item["source_herb_names"].add(row.get("herb_name", ""))
        item["compound_target_count"] += int(row.get("compound_target_count") or 0)
        item["pubchem_match_status"].add(row.get("pubchem_match_status", ""))
    output = []
    for item in grouped.values():
        output.append(
            {
                "ligand_id": item["ligand_id"],
                "pubchem_cid": item["pubchem_cid"],
                "inchikey": item["inchikey"],
                "canonical_smiles": item["canonical_smiles"],
                "isomeric_smiles": item["isomeric_smiles"],
                "source_ingredient_ids": "; ".join(sorted(value for value in item["source_ingredient_ids"] if value)),
                "source_ingredient_names": "; ".join(sorted(value for value in item["source_ingredient_names"] if value)),
                "source_herb_names": "; ".join(sorted(value for value in item["source_herb_names"] if value)),
                "ingredient_record_count": len([value for value in item["source_ingredient_ids"] if value]),
                "compound_target_count": item["compound_target_count"],
                "pubchem_match_status": "; ".join(sorted(value for value in item["pubchem_match_status"] if value)),
            }
        )
    output.sort(key=lambda row: row["ligand_id"])
    return output


def build_pubchem_unmatched_rows(active_rows: list[dict], target_count_by_mol: dict[str, int]) -> list[dict]:
    rows = []
    for row in active_rows:
        has_structure = bool(str(row.get("pubchem_cid", "")).strip() or str(row.get("smiles", "")).strip())
        if has_structure:
            continue
        rows.append(
            {
                "ingredient_id": row.get("ingredient_id", ""),
                "ingredient_name": row.get("ingredient_name", ""),
                "cleaned_name": row.get("normalized_ingredient_name", "") or normalize_pubchem_name(row.get("ingredient_name", "")),
                "attempted_names": row.get("pubchem_attempted_names", ""),
                "reason": row.get("pubchem_error", "") or row.get("pubchem_match_status", "") or "no_pubchem_structure",
                "used_in_network_analysis": "true",
                "eligible_for_docking": "false",
                "pubchem_error": row.get("pubchem_error", ""),
                "ob": row.get("ob", ""),
                "dl": row.get("dl", ""),
                "target_count": target_count_by_mol.get(row.get("ingredient_id", ""), 0),
            }
        )
    return rows


def write_module_outputs(
    *,
    params: dict,
    ingredient_dir: Path,
    target_dir: Path,
    pubchem_cache_path: Path,
    pubchem_cache: dict[str, dict[str, str]],
    shared_pubchem_cache_path: Path | None = None,
    raw_output: list[dict],
    active_rows: list[dict],
    compound_targets: list[dict],
    unmapped_targets: list[dict],
    ambiguous_targets: list[dict],
    ligand_ready_rows: list[dict],
    warnings: list[str],
    pubchem_enabled: bool,
    detail_url: str,
    pubchem_query_attempts: list[dict] | None = None,
    tcmsp_fetch_diagnostic: dict | None = None,
    source_merge_summary: dict | None = None,
    ambiguous_multi_name_rows: list[dict] | None = None,
    split_multi_name_rows: list[dict] | None = None,
) -> None:
    warnings = unique_preserve(warnings)
    input_section = params.get("input", {})
    input_mode = input_section.get("mode", "herb")
    herb_name = input_section.get("herb_name", "")
    compound_name = input_section.get("compound_name", "")
    diagnostic = tcmsp_fetch_diagnostic or {
        "requested_herb": herb_name or compound_name,
        "detail_url": detail_url,
        "http_status": "",
        "detected_reason": "not_applicable",
        "fallback_searched_paths": [],
        "final_decision": "used_remote_or_direct_tcmsp_record",
        "tcmsp_data_source": "remote",
        "tcmsp_source_conclusion": "Direct TCMSP record was used.",
        "tcmsp_fallback_used": False,
        "selected_cache_path": "",
        "generated_at": now_utc(),
    }

    target_count_by_mol: dict[str, int] = {}
    for row in compound_targets:
        target_count_by_mol[row["ingredient_id"]] = target_count_by_mol.get(row["ingredient_id"], 0) + 1
    pubchem_unmatched_rows = build_pubchem_unmatched_rows(active_rows, target_count_by_mol)
    deduplicated_active_rows = collapse_deduplicated_active_rows(active_rows)
    unique_ligand_ready_rows = collapse_unique_ligand_ready_rows(ligand_ready_rows)
    tcmsp_obdl_active_rows = [
        row
        for row in active_rows
        if "eligible_tcmsp_ob_dl" in str(row.get("network_eligibility_reason", ""))
    ]
    herbiv_candidate_rows = [
        row for row in active_rows
        if "HerbiV" in str(row.get("source_list", "") or row.get("source", ""))
        and str(row.get("herbiv_category", "")).lower() == "candidate"
    ]
    herbiv_exploratory_rows = [
        row for row in raw_output
        if "HerbiV" in str(row.get("source_list", "") or row.get("source", ""))
        and str(row.get("herbiv_category", "")).lower() == "exploratory"
    ]
    symmap_candidate_rows = [row for row in active_rows if "SymMap2" in str(row.get("source_list", "") or row.get("source", ""))]

    write_rows(ingredient_dir / "raw_ingredients.csv", raw_output, INGREDIENT_FIELDS)
    write_rows(ingredient_dir / "active_ingredients.csv", active_rows, INGREDIENT_FIELDS)
    write_rows(ingredient_dir / "tcmsp_obdl_active_ingredients.csv", tcmsp_obdl_active_rows, INGREDIENT_FIELDS)
    write_rows(ingredient_dir / "herbiv_candidate_ingredients.csv", herbiv_candidate_rows, INGREDIENT_FIELDS)
    write_rows(ingredient_dir / "herbiv_exploratory_ingredients.csv", herbiv_exploratory_rows, INGREDIENT_FIELDS)
    write_rows(ingredient_dir / "symmap_candidate_ingredients.csv", symmap_candidate_rows, INGREDIENT_FIELDS)
    write_rows(ingredient_dir / "merged_network_candidate_ingredients.csv", active_rows, INGREDIENT_FIELDS)
    write_rows(ingredient_dir / "deduplicated_active_ingredients.csv", deduplicated_active_rows, DEDUPLICATED_ACTIVE_FIELDS)
    write_rows(ingredient_dir / "ligand_ready_ingredients.csv", ligand_ready_rows, LIGAND_FIELDS)
    write_rows(ingredient_dir / "unique_ligand_ready_molecules.csv", unique_ligand_ready_rows, UNIQUE_LIGAND_FIELDS)
    write_rows(ingredient_dir / "pubchem_unmatched_ingredients.csv", pubchem_unmatched_rows, PUBCHEM_UNMATCHED_FIELDS)
    write_rows(ingredient_dir / "pubchem_query_attempts.csv", pubchem_query_attempts or [], PUBCHEM_QUERY_ATTEMPT_FIELDS)
    write_rows(ingredient_dir / "ambiguous_multi_name_ingredients.csv", ambiguous_multi_name_rows or [], MULTI_NAME_FIELDS)
    write_rows(ingredient_dir / "split_multi_name_ingredients.csv", split_multi_name_rows or [], MULTI_NAME_FIELDS)
    write_json(
        ingredient_dir / "ingredient_name_normalization_summary.json",
        {
            "generated_at": now_utc(),
            "ambiguous_multi_name_ingredients": len(ambiguous_multi_name_rows or []),
            "split_or_alias_multi_name_ingredients": len(split_multi_name_rows or []),
            "rule": "Comma/semicolon/slash separated names are excluded as ambiguous unless all parts normalize as aliases of the same compound.",
        },
    )
    structure_ready_rows = []
    for row in active_rows:
        if row.get("structure_status") in {"smiles_ready", "inchikey_and_smiles_ready"}:
            structure_ready_rows.append(
                {
                    **row,
                    "compound_target_count": target_count_by_mol.get(row.get("ingredient_id", ""), 0),
                    "structure_source": row.get("smiles_origin") or "PubChemCID",
                    "recommended_for_docking": "true" if target_count_by_mol.get(row.get("ingredient_id", ""), 0) > 0 else "false",
                }
            )
    docking_ready_rows = [
        row
        for row in structure_ready_rows
        if int(row.get("compound_target_count") or 0) > 0 and row.get("recommended_for_docking") == "true"
    ]
    cid_ready_rows = [row for row in active_rows if row.get("pubchem_cid")]
    write_rows(ingredient_dir / "structure_ready_ligands.csv", structure_ready_rows, LIGAND_FIELDS)
    write_rows(ingredient_dir / "docking_ready_ligands.csv", docking_ready_rows, LIGAND_FIELDS)
    write_json(
        ingredient_dir / "structure_readiness_summary.json",
        {
            "generated_at": now_utc(),
            "pubchem_cid_ready": len(cid_ready_rows),
            "smiles_ready": len(structure_ready_rows),
            "inchikey_ready": len([row for row in active_rows if row.get("inchikey")]),
            "structure_file_ready": 0,
            "docking_ready": len(docking_ready_rows),
            "docking_ready_definition": "SMILES/InChIKey-ready only; no PDBQT generation in this workflow version.",
        },
    )
    write_json(ingredient_dir / "tcmsp_fetch_diagnostic.json", diagnostic)
    write_rows(target_dir / "compound_targets.raw.csv", compound_targets, TARGET_FIELDS)
    write_rows(target_dir / "compound_targets.normalized.csv", compound_targets, TARGET_FIELDS)
    write_rows(target_dir / "unmapped_targets.csv", unmapped_targets, UNMAPPED_TARGET_FIELDS)
    write_rows(target_dir / "ambiguous_targets.csv", ambiguous_targets, AMBIGUOUS_TARGET_FIELDS)
    if source_merge_summary is None:
        source_merge_summary = {
            "generated_at": now_utc(),
            "status": "passed",
            "allowed_herb_compound_target_sources": ["TCMSP", "HerbiV", "SymMap2"],
            "source_rows": [
                {
                    "source": "TCMSP",
                    "status": "available",
                    "raw_ingredients": len(raw_output),
                    "raw_target_relations": len(compound_targets),
                    "source_detail": diagnostic.get("tcmsp_data_source", ""),
                }
            ],
            "supplemented_by_backup_sources": False,
        }
    final_target_relations_by_source = {"TCMSP": 0, "HerbiV": 0, "SymMap2": 0}
    for target_row in compound_targets:
        for source in _source_set(
            target_row.get("eligible_source_list")
            or target_row.get("source_list")
            or target_row.get("source")
        ):
            if source in final_target_relations_by_source:
                final_target_relations_by_source[source] += 1
    for source_row in source_merge_summary.get("source_rows", []):
        source = str(source_row.get("source", ""))
        if source in final_target_relations_by_source:
            source_row["final_main_network_target_relations"] = (
                final_target_relations_by_source[source]
            )
    source_merge_summary["final_main_network_target_relations_by_source"] = (
        final_target_relations_by_source
    )
    source_merge_summary["status"] = "passed" if compound_targets else "warning"
    source_merge_summary["source_coverage_status"] = (
        "covered" if compound_targets else "no_valid_source_coverage"
    )
    source_merge_summary["failure_reason"] = (
        ""
        if compound_targets
        else "no_main_network_edge_remained_after_final_source_specific_filters"
    )
    source_merge_summary["final_main_network_eligible_edges"] = len(compound_targets)
    source_merge_summary["supplemented_by_backup_sources"] = any(
        final_target_relations_by_source[source] > 0 for source in ("HerbiV", "SymMap2")
    )
    write_json(target_dir / "source_merge_summary.json", source_merge_summary)
    write_json(pubchem_cache_path, pubchem_cache)
    update_shared_pubchem_cache(shared_pubchem_cache_path, pubchem_cache)

    ingredient_names = unique_preserve(row["ingredient_name"] for row in active_rows)
    target_genes = unique_preserve(row["target_gene"] for row in compound_targets)
    active_with_smiles = [row for row in active_rows if str(row.get("smiles", "")).strip()]
    tcmsp_raw_rows = [
        row
        for row in raw_output
        if "TCMSP"
        in {
            part.strip()
            for part in str(row.get("source_list", "") or row.get("source", "")).split(";")
            if part.strip()
        }
    ]
    ob_dl_pass_count = 0
    target_mol_ids_for_qc = set(target_count_by_mol)
    for row in tcmsp_raw_rows:
        ob = parse_optional_float(row.get("ob"))
        dl = parse_optional_float(row.get("dl"))
        if ob is not None and dl is not None and ob >= float(params.get("compound_filter", {}).get("ob_min", 30)) and dl >= float(params.get("compound_filter", {}).get("dl_min", 0.18)):
            ob_dl_pass_count += 1
    no_target_filtered_count = sum(1 for row in raw_output if row.get("ingredient_id", "") not in target_mol_ids_for_qc)
    # OB/DL is a TCMSP-only policy. Supplemental-source rows must never enter
    # this denominator or be reported as failed TCMSP records.
    ob_dl_filtered_count = max(len(tcmsp_raw_rows) - ob_dl_pass_count, 0)
    source_qc_rows = {
        str(row.get("source", "")): row
        for row in (source_merge_summary or {}).get("source_rows", [])
    }
    herbiv_source_qc = source_qc_rows.get("HerbiV", {})
    symmap_source_qc = source_qc_rows.get("SymMap2", {})
    smiles_filtered_count = 0
    active_without_smiles_count = len(active_rows) - len(active_with_smiles)
    mapping_unmapped = sum(int(row.get("target_row_count") or 0) for row in unmapped_targets if row.get("mapping_status") == "unmapped")
    mapping_ignored = sum(int(row.get("target_row_count") or 0) for row in unmapped_targets if row.get("mapping_status") == "ignored_non_human_or_irrelevant")
    mapping_ambiguous = sum(int(row.get("target_row_count") or 0) for row in ambiguous_targets)
    mapping_total = len(compound_targets) + mapping_unmapped + mapping_ignored + mapping_ambiguous
    mapping_success = len(compound_targets)
    mapping_success_rate = (mapping_success / mapping_total) if mapping_total else 0
    summary_warnings = unique_preserve(warnings)
    if diagnostic.get("completeness_warning"):
        summary_warnings.append(str(diagnostic.get("completeness_warning")))
    if diagnostic.get("herb_match_warning"):
        summary_warnings.append(str(diagnostic.get("herb_match_warning")))
    if diagnostic.get("tcmsp_fallback_used"):
        summary_warnings.append(
            f"TCMSP remote herb detail was unavailable; used {diagnostic.get('tcmsp_data_source')} from {diagnostic.get('selected_cache_path')}."
        )
    if pubchem_unmatched_rows:
        summary_warnings.append(
            f"PubChem structure lookup did not match {len(pubchem_unmatched_rows)} active ingredients; see 01_ingredients/pubchem_unmatched_ingredients.csv."
        )
    if mapping_success_rate < 0.70 and mapping_total:
        summary_warnings.append(
            f"TCMSP target mapping success rate is {mapping_success_rate:.1%}; see 02_compound_targets/target_mapping_summary.json."
        )
    if mapping_unmapped:
        summary_warnings.append(
            f"{mapping_unmapped} TCMSP target rows remain unmapped; see 02_compound_targets/unmapped_targets.csv."
        )
    if mapping_ambiguous:
        summary_warnings.append(
            f"{mapping_ambiguous} TCMSP target rows are ambiguous; see 02_compound_targets/ambiguous_targets.csv."
        )

    generated_ingredient_files = [
        str(ingredient_dir / "raw_ingredients.csv"),
        str(ingredient_dir / "active_ingredients.csv"),
        str(ingredient_dir / "tcmsp_obdl_active_ingredients.csv"),
        str(ingredient_dir / "herbiv_candidate_ingredients.csv"),
        str(ingredient_dir / "herbiv_exploratory_ingredients.csv"),
        str(ingredient_dir / "symmap_candidate_ingredients.csv"),
        str(ingredient_dir / "merged_network_candidate_ingredients.csv"),
        str(ingredient_dir / "deduplicated_active_ingredients.csv"),
        str(ingredient_dir / "ligand_ready_ingredients.csv"),
        str(ingredient_dir / "unique_ligand_ready_molecules.csv"),
        str(ingredient_dir / "pubchem_unmatched_ingredients.csv"),
        str(ingredient_dir / "pubchem_query_attempts.csv"),
        str(ingredient_dir / "ambiguous_multi_name_ingredients.csv"),
        str(ingredient_dir / "split_multi_name_ingredients.csv"),
        str(ingredient_dir / "ingredient_name_normalization_summary.json"),
        str(ingredient_dir / "structure_readiness_summary.json"),
        str(ingredient_dir / "union_ingredients.csv"),
        str(ingredient_dir / "tcmsp_fetch_diagnostic.json"),
    ]
    generated_target_files = [
        str(target_dir / "compound_targets.raw.csv"),
        str(target_dir / "compound_targets.normalized.csv"),
        str(target_dir / "source_merge_summary.json"),
        str(target_dir / "source_reliability_audit.json"),
        str(target_dir / "supplemental_source_audit.json"),
        str(target_dir / "herbiv_source_audit.json"),
        str(target_dir / "symmap_source_audit.json"),
        str(target_dir / "union_targets.csv"),
        str(target_dir / "union_ingredient_target_edges.csv"),
        str(target_dir / "unmapped_targets.csv"),
        str(target_dir / "ambiguous_targets.csv"),
        str(target_dir / "target_mapping_summary.json"),
    ]
    common_summary = {
        "input_mode": input_mode,
        "herb_name": herb_name,
        "compound_name": compound_name,
        "herb_name_standardization": diagnostic.get("herb_name_standardization", {}),
        "original_input_herb_name": diagnostic.get("original_input_herb_name", herb_name),
        "standardized_herb_for_lookup": diagnostic.get("standardized_herb_for_lookup", herb_name),
        "source_coverage_status": (source_merge_summary or {}).get("source_coverage_status", ""),
        "detail_url": detail_url,
        "tcmsp_data_source": diagnostic.get("tcmsp_data_source", "remote"),
        "tcmsp_source_conclusion": diagnostic.get("tcmsp_source_conclusion", ""),
        "tcmsp_fallback_used": bool(diagnostic.get("tcmsp_fallback_used", False)),
        "remote_attempted": bool(diagnostic.get("remote_attempted", False)),
        "remote_success": bool(diagnostic.get("remote_success", False)),
        "remote_failure_reason": diagnostic.get("remote_failure_reason", ""),
        "remote_detail_url": diagnostic.get("remote_detail_url", ""),
        "remote_http_status": diagnostic.get("remote_http_status", ""),
        "remote_response_length": diagnostic.get("remote_response_length", ""),
        "remote_primary_search_url": diagnostic.get("remote_primary_search_url", ""),
        "remote_search_url": diagnostic.get("remote_search_url", ""),
        "remote_last_attempt_url": diagnostic.get("remote_last_attempt_url", ""),
        "remote_selected_candidate_url": diagnostic.get("remote_selected_candidate_url", ""),
        "remote_selected_candidate_text": diagnostic.get("remote_selected_candidate_text", ""),
        "remote_selected_candidate_match_field": diagnostic.get("remote_selected_candidate_match_field", ""),
        "remote_selected_candidate_match_score": diagnostic.get("remote_selected_candidate_match_score", ""),
        "remote_candidate_count": diagnostic.get("remote_candidate_count", 0),
        "remote_candidates": diagnostic.get("remote_candidates", []),
        "remote_candidate_failure_reason": diagnostic.get("remote_candidate_failure_reason", ""),
        "remote_page_type": diagnostic.get("remote_page_type", ""),
        "remote_http_ok": bool(diagnostic.get("remote_http_ok", False)),
        "remote_html_valid": bool(diagnostic.get("remote_html_valid", False)),
        "remote_grid_found": bool(diagnostic.get("remote_grid_found", False)),
        "remote_grid2_found": bool(diagnostic.get("remote_grid2_found", False)),
        "remote_grid_row_count": int(diagnostic.get("remote_grid_row_count", 0) or 0),
        "remote_grid2_row_count": int(diagnostic.get("remote_grid2_row_count", 0) or 0),
        "remote_grid_required_fields_valid": bool(diagnostic.get("remote_grid_required_fields_valid", False)),
        "remote_grid2_required_fields_valid": bool(diagnostic.get("remote_grid2_required_fields_valid", False)),
        "remote_parse_success": bool(diagnostic.get("remote_parse_success", False)),
        "live_tcmsp_scrape_validated": bool(diagnostic.get("live_tcmsp_scrape_validated", False)),
        "fallback_snapshot_validated": bool(diagnostic.get("fallback_snapshot_validated", False)),
        "final_tcmsp_source_is_live": bool(diagnostic.get("final_tcmsp_source_is_live", False)),
        "actual_source_attempt_order": diagnostic.get("actual_source_attempt_order", []),
        "remote_search_attempts": diagnostic.get("remote_search_attempts", []),
        "fallback_after_remote": bool(diagnostic.get("fallback_after_remote", False)),
        "mapping_total": mapping_total,
        "mapping_success": mapping_success,
        "mapping_failed": mapping_unmapped,
        "mapping_ambiguous": mapping_ambiguous,
        "mapping_ignored": mapping_ignored,
        "mapping_success_rate": round(mapping_success_rate, 6),
        "target_mapping_total": mapping_total,
        "target_mapping_success": mapping_success,
        "target_mapping_failed": mapping_unmapped,
        "target_mapping_ambiguous": mapping_ambiguous,
        "target_mapping_ignored": mapping_ignored,
        "target_mapping_success_rate": round(mapping_success_rate, 6),
        "official_download_method": diagnostic.get("official_download_method", ""),
        "download_file_paths": diagnostic.get("download_file_paths", []),
        "input_herb_name": diagnostic.get("input_herb_name", herb_name),
        "herb_match_field": diagnostic.get("herb_match_field", ""),
        "matched_herb_id": diagnostic.get("matched_herb_id", ""),
        "matched_herb_name": diagnostic.get("matched_herb_name", ""),
        "herb_match_score": diagnostic.get("herb_match_score", ""),
        "herb_match_type": diagnostic.get("herb_match_type", ""),
        "herb_exact_match": bool(diagnostic.get("herb_exact_match", False)),
        "herb_alias_match": bool(diagnostic.get("herb_alias_match", False)),
        "herb_fuzzy_match_used": bool(diagnostic.get("herb_fuzzy_match_used", False)),
        "herb_match_requires_user_confirmation": bool(diagnostic.get("herb_match_requires_user_confirmation", False)),
        "herb_match_warning": diagnostic.get("herb_match_warning", ""),
        "herb_match_status": diagnostic.get("herb_match_status", ""),
        "herb_match_queries": diagnostic.get("herb_match_queries", []),
        "aliases_used": diagnostic.get("aliases_used", []),
        "source_detail": diagnostic.get("source_detail", ""),
        "source_urls": diagnostic.get("source_urls", []),
        "snapshot_manifest_path": diagnostic.get("snapshot_manifest_path", ""),
        "snapshot_db_path": diagnostic.get("snapshot_db_path", ""),
        "snapshot_source_type": diagnostic.get("snapshot_source_type", ""),
        "snapshot_row_counts": diagnostic.get("snapshot_row_counts", {}),
        "snapshot_audit_path": diagnostic.get("snapshot_audit_path", ""),
        "is_full_snapshot": bool(diagnostic.get("is_full_snapshot", False)),
        "is_seed_cache": bool(diagnostic.get("is_seed_cache", False)),
        "is_official_download_files": bool(diagnostic.get("is_official_download_files", False)),
        "source_completeness_level": diagnostic.get("source_completeness_level", ""),
        "completeness_warning": diagnostic.get("completeness_warning", ""),
        "raw_target_relation_count": source_merge_summary.get(
            "raw_target_relations_before_merge", diagnostic.get("raw_target_relation_count", len(compound_targets))
        ),
        "normalized_target_count": len(compound_targets),
        "source_merge_summary": source_merge_summary,
        "supplemented_by_backup_sources": bool(source_merge_summary.get("supplemented_by_backup_sources", False)),
    }
    tcmsp_data_quality = {
        "data_source": common_summary["tcmsp_data_source"],
        "tcmsp_source_conclusion": common_summary["tcmsp_source_conclusion"],
        "remote_attempted": common_summary["remote_attempted"],
        "remote_success": common_summary["remote_success"],
        "remote_failure_reason": common_summary["remote_failure_reason"],
        "remote_detail_url": common_summary["remote_detail_url"],
        "remote_http_status": common_summary["remote_http_status"],
        "remote_response_length": common_summary["remote_response_length"],
        "remote_primary_search_url": common_summary["remote_primary_search_url"],
        "remote_search_url": common_summary["remote_search_url"],
        "remote_last_attempt_url": common_summary["remote_last_attempt_url"],
        "remote_selected_candidate_url": common_summary["remote_selected_candidate_url"],
        "remote_selected_candidate_text": common_summary["remote_selected_candidate_text"],
        "remote_selected_candidate_match_field": common_summary["remote_selected_candidate_match_field"],
        "remote_selected_candidate_match_score": common_summary["remote_selected_candidate_match_score"],
        "remote_candidate_count": common_summary["remote_candidate_count"],
        "remote_candidates": common_summary["remote_candidates"],
        "remote_candidate_failure_reason": common_summary["remote_candidate_failure_reason"],
        "remote_page_type": common_summary["remote_page_type"],
        "remote_http_ok": common_summary["remote_http_ok"],
        "remote_html_valid": common_summary["remote_html_valid"],
        "remote_grid_found": common_summary["remote_grid_found"],
        "remote_grid2_found": common_summary["remote_grid2_found"],
        "remote_grid_row_count": common_summary["remote_grid_row_count"],
        "remote_grid2_row_count": common_summary["remote_grid2_row_count"],
        "remote_grid_required_fields_valid": common_summary["remote_grid_required_fields_valid"],
        "remote_grid2_required_fields_valid": common_summary["remote_grid2_required_fields_valid"],
        "remote_parse_success": common_summary["remote_parse_success"],
        "live_tcmsp_scrape_validated": common_summary["live_tcmsp_scrape_validated"],
        "fallback_snapshot_validated": common_summary["fallback_snapshot_validated"],
        "final_tcmsp_source_is_live": common_summary["final_tcmsp_source_is_live"],
        "actual_source_attempt_order": common_summary["actual_source_attempt_order"],
        "remote_search_attempts": common_summary["remote_search_attempts"],
        "fallback_after_remote": common_summary["fallback_after_remote"],
        "source_detail": common_summary["source_detail"],
        "source_completeness_level": common_summary["source_completeness_level"],
        "is_full_snapshot": common_summary["is_full_snapshot"],
        "is_seed_cache": common_summary["is_seed_cache"],
        "is_official_download_files": common_summary["is_official_download_files"],
        "snapshot_db_path": common_summary["snapshot_db_path"],
        "snapshot_manifest_path": common_summary["snapshot_manifest_path"],
        "snapshot_audit_path": common_summary["snapshot_audit_path"],
        "snapshot_source_type": common_summary["snapshot_source_type"],
        "snapshot_row_counts": common_summary["snapshot_row_counts"],
        "source_files": common_summary["download_file_paths"],
        "source_urls": common_summary["source_urls"],
        "input_herb_name": common_summary["input_herb_name"],
        "matched_herb_id": common_summary["matched_herb_id"],
        "matched_herb_name": common_summary["matched_herb_name"],
        "herb_match_field": common_summary["herb_match_field"],
        "herb_match_score": common_summary["herb_match_score"],
        "herb_match_type": common_summary["herb_match_type"],
        "herb_exact_match": common_summary["herb_exact_match"],
        "herb_alias_match": common_summary["herb_alias_match"],
        "herb_fuzzy_match_used": common_summary["herb_fuzzy_match_used"],
        "herb_match_requires_user_confirmation": common_summary["herb_match_requires_user_confirmation"],
        "herb_match_warning": common_summary["herb_match_warning"],
        "herb_match_status": common_summary["herb_match_status"],
        "herb_match_queries": common_summary["herb_match_queries"],
        "aliases_used": common_summary["aliases_used"],
        "raw_ingredient_count": len(raw_output),
        "ob_dl_pass_ingredient_count": ob_dl_pass_count,
        "active_ingredient_count": len(active_rows),
        "raw_target_relation_count": common_summary["raw_target_relation_count"],
        "normalized_target_count": len(compound_targets),
        "completeness_warning": common_summary["completeness_warning"],
        "source_merge_summary": source_merge_summary,
        "supplemented_by_backup_sources": common_summary["supplemented_by_backup_sources"],
    }
    write_json(
        target_dir / "target_mapping_summary.json",
        {
            "module": "fetch_herb_compounds",
            "status": "passed" if mapping_success_rate >= 0.70 or mapping_total == 0 else "warning",
            "generated_at": now_utc(),
            **common_summary,
            "tcmsp_data_quality": tcmsp_data_quality,
            "unique_mapped_genes": len(target_genes),
            "unmapped_target_names": len([row for row in unmapped_targets if row.get("mapping_status") == "unmapped"]),
            "ambiguous_target_names": len(ambiguous_targets),
            "ignored_target_names": len([row for row in unmapped_targets if row.get("mapping_status") == "ignored_non_human_or_irrelevant"]),
            "generated_files": generated_target_files,
            "warnings": summary_warnings[:100],
        },
    )
    write_json(
        ingredient_dir / "ingredient_qc.json",
        {
            "module": "fetch_herb_compounds",
            "status": "passed" if active_rows and compound_targets else "warning",
            "generated_at": now_utc(),
            "source": "TCMSP;HerbiV;SymMap2",
            "primary_herb_compound_target_source": "TCMSP",
            "supplemental_herb_compound_target_sources": ["HerbiV", "SymMap2"],
            **common_summary,
            "tcmsp_data_quality": tcmsp_data_quality,
            "raw_ingredient_count": len(raw_output),
            "tcmsp_ob_dl_evaluated_ingredient_count": len(tcmsp_raw_rows),
            "herbiv_raw_ingredient_count": int(herbiv_source_qc.get("raw_ingredients") or 0),
            "herbiv_raw_target_relation_count": int(herbiv_source_qc.get("raw_target_relations") or 0),
            "symmap_raw_ingredient_count": int(symmap_source_qc.get("raw_ingredients") or 0),
            "symmap_raw_target_relation_count": int(symmap_source_qc.get("raw_target_relations") or 0),
            "active_ingredient_records": len(active_rows),
            "tcmsp_obdl_active_ingredients": len(tcmsp_obdl_active_rows),
            "herbiv_candidate_ingredients": len(herbiv_candidate_rows),
            "herbiv_exploratory_ingredients": len(herbiv_exploratory_rows),
            "symmap_candidate_ingredients": len(symmap_candidate_rows),
            "merged_network_candidate_ingredients": len(active_rows),
            "deduplicated_active_ingredients": len(deduplicated_active_rows),
            "active_ingredient_count": len(active_rows),
            "active_ingredients_total": len(active_rows),
            "active_ingredients_with_smiles": len(active_with_smiles),
            "active_ingredients_without_smiles": active_without_smiles_count,
            "ob_dl_pass_ingredient_count": ob_dl_pass_count,
            "ob_dl_filtered_ingredient_count": ob_dl_filtered_count,
            "no_target_filtered_ingredient_count": no_target_filtered_count,
            "smiles_filtered_ingredient_count": smiles_filtered_count,
            "active_ingredient_count_explanation": (
                f"Merged raw ingredients={len(raw_output)}; TCMSP OB/DL evaluated={len(tcmsp_raw_rows)}; "
                f"TCMSP OB/DL pass={ob_dl_pass_count}; TCMSP filtered_by_ob_dl={ob_dl_filtered_count}; "
                f"filtered_by_no_target={no_target_filtered_count}; "
                f"filtered_by_smiles={smiles_filtered_count}; merged_network_candidates={len(active_rows)}. "
                f"HerbiV raw ingredients={int(herbiv_source_qc.get('raw_ingredients') or 0)} and "
                f"SymMap2 raw ingredients={int(symmap_source_qc.get('raw_ingredients') or 0)} are reported separately. "
                "TCMSP records alone are OB/DL-filtered; HerbiV/SymMap2 records follow their source policies. "
                "SMILES absence is tracked for docking readiness and does not remove network candidates."
            ),
            "ligand_ready_ingredients": len(ligand_ready_rows),
            "ligand_ready_ingredient_records": len(ligand_ready_rows),
            "unique_ligand_ready_molecules": len(unique_ligand_ready_rows),
            "pubchem_cid_ready_ligands": len(cid_ready_rows),
            "smiles_ready_ligands": len(structure_ready_rows),
            "structure_file_ready_ligands": 0,
            "docking_ready_ligands": len(docking_ready_rows),
            "ambiguous_multi_name_ingredients": len(ambiguous_multi_name_rows or []),
            "split_multi_name_ingredients": len(split_multi_name_rows or []),
            "pubchem_unmatched_ingredients": len(pubchem_unmatched_rows),
            "pubchem_query_attempts": len(pubchem_query_attempts or []),
            "compound_target_count": len(compound_targets),
            "unique_target_count": len(target_genes),
            "pubchem_enabled": pubchem_enabled,
            "active_ingredients_preview": ingredient_names[:20],
            "generated_files": generated_ingredient_files,
            "warnings": summary_warnings[:100],
        },
    )
    write_json(
        target_dir / "compound_target_summary.json",
        {
            "module": "fetch_herb_compounds",
            "status": "passed" if compound_targets else "warning",
            "generated_at": now_utc(),
            "compound_target_count": len(compound_targets),
            "unique_target_count": len(target_genes),
            **common_summary,
            "tcmsp_data_quality": tcmsp_data_quality,
            "generated_files": generated_target_files,
            "warnings": summary_warnings[:100],
        },
    )


def run_compound_mode(
    *,
    params: dict,
    args: argparse.Namespace,
    ingredient_dir: Path,
    target_dir: Path,
    pubchem_cache_path: Path,
    pubchem_cache: dict[str, dict[str, str]],
    shared_pubchem_cache_path: Path | None = None,
) -> None:
    requested_name = params.get("input", {}).get("compound_name", "")
    compound_name = normalize_compound_input(requested_name)
    if not compound_name:
        fail("input.compound_name is required in compound mode.")

    record = find_molecule_record(compound_name)
    molecule_id = str(record.get("molecule_ID", "") or record.get("MOL_ID", "") or "")
    if not molecule_id:
        fail(f"TCMSP molecule search did not provide molecule_ID for {compound_name}.")
    molecule_page_id = molecule_id[3:] if molecule_id.upper().startswith("MOL") else molecule_id
    detail_url = f"{TCMSP_BASE}/molecule.php?qn={quote(molecule_page_id)}"
    response = http_get_tcmsp(detail_url, timeout=60)
    if response.status_code != 200:
        fail(f"TCMSP molecule detail returned HTTP {response.status_code}: {detail_url}")

    mol_rows = extract_js_var_array(response.text, "mol")
    target_rows = extract_js_var_array(response.text, "tar")
    mol = mol_rows[0] if mol_rows else record
    mol_id = str(mol.get("MOL_ID", "") or record.get("MOL_ID", "") or molecule_id)
    raw_ingredient_name = str(mol.get("molecule_name", "") or record.get("molecule_name", "") or compound_name)
    name_info = normalize_ingredient_display_name(raw_ingredient_name)
    ingredient_name = name_info["ingredient_name"]
    pubchem_cid = str(mol.get("PubChem_ID", "") or mol.get("pubchem_cid", "") or "")
    smiles = str(mol.get("smiles", "") or "")
    inchikey = str(mol.get("InChIKey", "") or mol.get("inchi_key", "") or "")
    pubchem_query_attempts: list[dict] = []
    structure = (
        fetch_pubchem_structure(
            ingredient_name,
            pubchem_cache,
            pubchem_query_attempts,
            ingredient_id=mol_id,
            raw_ingredient_name=raw_ingredient_name,
        )
        if args.pubchem
        else empty_pubchem_result()
    )
    if pubchem_cid and not structure.get("pubchem_cid"):
        structure["pubchem_cid"] = pubchem_cid
    if smiles and not structure.get("smiles"):
        structure["smiles"] = smiles
        structure["canonical_smiles"] = smiles
    if inchikey and not structure.get("inchikey"):
        structure["inchikey"] = inchikey
        structure["inchi_key"] = inchikey
    if structure.get("pubchem_match_status") == "not_queried" and (pubchem_cid or smiles or inchikey):
        structure["pubchem_match_status"] = "matched_tcmsp"
    structure_status, docking_ready_reason, smiles_origin, inchikey_origin = structure_status_for_row(
        structure,
        source_pubchem_cid=pubchem_cid,
    )

    output_row = {
        "herb_name": "",
        "ingredient_id": mol_id,
        "raw_ingredient_name": raw_ingredient_name,
        "ingredient_name": ingredient_name,
        "normalized_ingredient_name": name_info["normalized_ingredient_name"],
        "name_repair_applied": name_info["name_repair_applied"],
        "pubchem_cid": structure.get("pubchem_cid", ""),
        "smiles": structure.get("smiles", ""),
        "canonical_smiles": structure.get("canonical_smiles", ""),
        "isomeric_smiles": structure.get("isomeric_smiles", ""),
        "inchikey": structure.get("inchikey", ""),
        "inchi_key": structure.get("inchi_key", ""),
        "pubchem_match_status": structure.get("pubchem_match_status", ""),
        "pubchem_query_name": structure.get("pubchem_query_name", ""),
        "pubchem_attempted_names": structure.get("pubchem_attempted_names", ""),
        "pubchem_error": structure.get("pubchem_error", ""),
        "cid_origin": "source_database_cid" if pubchem_cid else "",
        "cid_origin_source_database": "TCMSP" if pubchem_cid else "",
        "pubchem_query_attempted": "true" if args.pubchem else "false",
        "pubchem_query_status": "matched_structure_by_name_or_source" if structure_status in {"smiles_ready", "inchikey_and_smiles_ready"} else structure.get("pubchem_match_status", ""),
        "smiles_origin": smiles_origin or ("TCMSP" if smiles else ""),
        "inchikey_origin": inchikey_origin or ("TCMSP" if inchikey else ""),
        "structure_file_path": "",
        "structure_status": structure_status,
        "docking_ready_reason": docking_ready_reason,
        "ob": mol.get("ob", ""),
        "dl": mol.get("dl", ""),
        "source": "TCMSP",
        "tcmsp_data_source": "tcmsp_molecule_page",
        "passed_filter": "true",
    }

    target_cache: dict[str, dict[str, str]] = {}
    compound_targets: list[dict] = []
    unmapped_source_rows: list[dict] = []
    ambiguous_source_rows: list[dict] = []
    warnings: list[str] = []
    for row in target_rows:
        target_name = row.get("target_name", "")
        mapping = map_target_record(target_name, target_cache) if target_name else make_mapping_record("", failure_reason="empty_target_name")
        gene = mapping.get("target_gene", "")
        if not gene:
            issue_row = {
                **mapping,
                "ingredient_id": mol_id,
                "raw_ingredient_name": raw_ingredient_name,
                "ingredient_name": ingredient_name,
            }
            if mapping.get("mapping_status") == "ambiguous":
                ambiguous_source_rows.append(issue_row)
            else:
                unmapped_source_rows.append(issue_row)
            continue
        raw_evidence_type = str(
            row.get("target_evidence_types")
            or row.get("evidence_type")
            or row.get("validated")
            or ""
        ).strip()
        compound_targets.append(
            {
                "herb_name": "",
                "raw_ingredient_name": raw_ingredient_name,
                "ingredient_name": ingredient_name,
                "normalized_ingredient_name": name_info["normalized_ingredient_name"],
                "ingredient_id": mol_id,
                "pubchem_cid": output_row.get("pubchem_cid", ""),
                "smiles": output_row.get("smiles", ""),
                "canonical_smiles": output_row.get("canonical_smiles", ""),
                "isomeric_smiles": output_row.get("isomeric_smiles", ""),
                "inchikey": output_row.get("inchikey", ""),
                "target_gene": gene,
                "uniprot_id": mapping.get("uniprot_id", ""),
                "target_name": target_name,
                "mapping_status": mapping.get("mapping_status", ""),
                "mapping_note": mapping.get("mapping_note", ""),
                "evidence_type": raw_evidence_type,
                "evidence_class": classify_target_evidence(raw_evidence_type),
                "mapping_validated": "true"
                if mapping.get("mapping_status") == "mapped_uniprot"
                else "false",
                "score": "",
                "source": "TCMSP",
                "source_list": "TCMSP",
                "eligible_source_list": "TCMSP",
                "all_source_list": "TCMSP",
                "target_evidence_types": raw_evidence_type,
                "network_eligibility": "eligible",
                "network_eligibility_reason": "direct_compound_mode_target_relation",
                "tcmsp_data_source": "tcmsp_molecule_page",
            }
        )

    write_module_outputs(
        params=params,
        ingredient_dir=ingredient_dir,
        target_dir=target_dir,
        pubchem_cache_path=pubchem_cache_path,
        pubchem_cache=pubchem_cache,
        shared_pubchem_cache_path=shared_pubchem_cache_path,
        raw_output=[output_row],
        active_rows=[output_row],
        compound_targets=compound_targets,
        unmapped_targets=collect_unmapped_targets(unmapped_source_rows),
        ambiguous_targets=collect_mapping_issues(ambiguous_source_rows),
        ligand_ready_rows=build_ligand_ready_rows([output_row], compound_targets),
        warnings=warnings,
        pubchem_enabled=args.pubchem,
        detail_url=detail_url,
        pubchem_query_attempts=pubchem_query_attempts,
    )
    print(f"Fetched compound {ingredient_name} and {len(compound_targets)} compound-target rows.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Fetch TCMSP herb compounds and compound targets.")
    parser.add_argument("--config", default="results/network_pharmacology/00_input/params.resolved.yaml")
    parser.add_argument("--pubchem", action="store_true", help="Enable PubChem structure lookup.")
    parser.add_argument("--uniprot", action="store_true", help="Enable UniProt fallback mapping for TCMSP target names.")
    args = parser.parse_args()

    project_root = project_root_from_cwd()
    params = load_params(args.config)
    out_root = output_root(params)
    ingredient_dir = ensure_dir(out_root / "01_ingredients")
    target_dir = ensure_dir(out_root / "02_compound_targets")
    cache_dir = ensure_dir(out_root / "_cache")
    pubchem_cache_path = cache_dir / "pubchem_structures.json"
    pubchem_cache = load_pubchem_cache(pubchem_cache_path)
    shared_pubchem_cache_path = project_root / "data" / "pubchem_structure_cache.json"
    merge_pubchem_cache(pubchem_cache, load_pubchem_cache(shared_pubchem_cache_path))

    input_mode = str(params.get("input", {}).get("mode", "herb")).strip().lower() or "herb"
    if input_mode == "compound":
        run_compound_mode(
            params=params,
            args=args,
            ingredient_dir=ingredient_dir,
            target_dir=target_dir,
            pubchem_cache_path=pubchem_cache_path,
            pubchem_cache=pubchem_cache,
            shared_pubchem_cache_path=shared_pubchem_cache_path,
        )
        return

    herb_name = params["input"]["herb_name"]
    herb_standardization = standardize_herb_name(herb_name)
    lookup_herb_name = herb_standardization["standardized_herb_for_lookup"]
    lookup_aliases = herb_standardization["aliases_for_lookup"]
    input_dir = ensure_dir(out_root / "00_input")
    write_json(input_dir / "herb_name_standardization.json", herb_standardization)
    filters = params.get("compound_filter", {})

    ingredient_rows, target_rows, detail_url, html, tcmsp_diagnostic = fetch_herb_detail_data(
        project_root=project_root,
        out_root=out_root,
        cache_dir=cache_dir,
        herb_name=lookup_herb_name,
    )
    tcmsp_diagnostic["original_input_herb_name"] = herb_name
    tcmsp_diagnostic["standardized_herb_for_lookup"] = lookup_herb_name
    tcmsp_diagnostic["herb_name_standardization"] = herb_standardization
    if herb_standardization.get("requires_user_confirmation"):
        tcmsp_diagnostic["input_herb_name"] = herb_name
        tcmsp_diagnostic["resolved_herb_name"] = lookup_herb_name
        tcmsp_diagnostic["canonical_herb_name"] = lookup_herb_name
        tcmsp_diagnostic["herb_match_requires_user_confirmation"] = True
        tcmsp_diagnostic["herb_match_warning"] = herb_standardization.get("warning", "")
        tcmsp_diagnostic["herb_match_status"] = herb_standardization.get("herb_match_status", "alias_needs_confirmation")
        if not tcmsp_diagnostic.get("herb_match_type"):
            tcmsp_diagnostic["herb_match_type"] = "ambiguous_alias"
        tcmsp_diagnostic["aliases_used"] = lookup_aliases
        write_json(ingredient_dir / "tcmsp_fetch_diagnostic.json", tcmsp_diagnostic)
        write_tcmsp_debug_meta(out_root, tcmsp_diagnostic)
    supplemental_results = fetch_supplemental_herb_sources(
        project_root=project_root,
        herb_name=lookup_herb_name,
        aliases=lookup_aliases,
        out_root=out_root,
        source_weights=params.get("source_weights", {}),
    )
    ingredient_rows, target_rows, source_merge_summary = merge_herb_source_rows(
        ingredient_rows,
        target_rows,
        supplemental_results,
        tcmsp_source_name=tcmsp_diagnostic.get("tcmsp_data_source", ""),
        out_root=out_root,
    )
    target_mol_ids = {row.get("MOL_ID", "") for row in target_rows}

    active_by_mol: dict[str, dict] = {}
    raw_output: list[dict] = []
    warnings: list[str] = []
    pubchem_query_attempts: list[dict] = []
    ambiguous_multi_name_rows: list[dict] = []
    split_multi_name_rows: list[dict] = []
    row_tcmsp_data_source = tcmsp_diagnostic.get("tcmsp_data_source", "")
    source_herb_name = (
        tcmsp_diagnostic.get("matched_herb_name")
        or tcmsp_diagnostic.get("resolved_herb_name")
        or lookup_herb_name
        or herb_name
    )

    for row in ingredient_rows:
        mol_id = row.get("MOL_ID", "")
        raw_name = row.get("molecule_name", "")
        name_info = normalize_ingredient_display_name(raw_name)
        name = name_info["ingredient_name"]
        passed, eligibility_reason = ingredient_eligibility(row, target_mol_ids, filters)
        source_pubchem_text = str(row.get("PubChem_ID", "") or row.get("pubchem_cid", "") or "").strip()
        source_pubchem_match = re.search(r"\d+", source_pubchem_text)
        source_pubchem_cid = source_pubchem_match.group(0) if source_pubchem_match else ""
        source_list_for_row = row.get("source_list", "") or row.get("source", "") or "TCMSP"
        multi_class, split_names, multi_reason = classify_multi_name(raw_name)
        if multi_class in {"ambiguous_multi_name", "alias_multi_name"}:
            multi_row = {
                "ingredient_id": mol_id,
                "raw_ingredient_name": raw_name,
                "ingredient_name": name,
                "normalized_ingredient_name": name_info["normalized_ingredient_name"],
                "source_list": source_list_for_row,
                "classification": multi_class,
                "reason": multi_reason,
                "split_names": "; ".join(split_names),
            }
            if multi_class == "ambiguous_multi_name":
                ambiguous_multi_name_rows.append(multi_row)
                passed = False
            else:
                split_multi_name_rows.append(multi_row)

        if passed and source_pubchem_cid and args.pubchem:
            structure = fetch_pubchem_structure_by_cid(
                source_pubchem_cid,
                pubchem_cache,
                pubchem_query_attempts,
                ingredient_id=mol_id,
                ingredient_name=name,
                raw_ingredient_name=raw_name,
                cid_origin_source=source_list_for_row,
            )
        elif passed and source_pubchem_cid:
            structure = empty_pubchem_result()
            structure.update(
                {
                    "pubchem_cid": source_pubchem_cid,
                    "pubchem_match_status": "source_pubchem_cid_not_queried",
                    "pubchem_query_name": name,
                    "pubchem_attempted_names": source_pubchem_cid,
                    "pubchem_error": "PubChem CID supplied by source database; property lookup disabled.",
                }
            )
        elif multi_class == "ambiguous_multi_name":
            structure = empty_pubchem_result()
            structure.update(
                {
                    "pubchem_match_status": "skipped_ambiguous_multi_name",
                    "pubchem_query_name": name,
                    "pubchem_attempted_names": raw_name,
                    "pubchem_error": "ambiguous TCMSP multi-name record; no single PubChem structure assigned",
                }
            )
        elif len(name) > 100:
            structure = empty_pubchem_result()
            structure.update(
                {
                    "pubchem_match_status": "skipped_long_complex_name",
                    "pubchem_query_name": name,
                    "pubchem_attempted_names": raw_name,
                    "pubchem_error": "long complex name skipped to avoid low-confidence PubChem matching",
                }
            )
        else:
            # 无 source CID — 如果 pubchem 已启用，按名称查 PubChem
            if passed and args.pubchem:
                structure = fetch_pubchem_structure(
                    name,
                    pubchem_cache,
                    pubchem_query_attempts,
                    ingredient_id=mol_id,
                    raw_ingredient_name=raw_name,
                )
            else:
                structure = empty_pubchem_result()
                structure.update(
                    {
                        "pubchem_match_status": "not_attempted_no_source_cid",
                        "pubchem_query_name": name,
                        "pubchem_attempted_names": "",
                        "pubchem_error": "no source PubChem CID; remote PubChem lookup skipped for stable batch execution",
                    }
                )
        structure_status, docking_ready_reason, smiles_origin, inchikey_origin = structure_status_for_row(
            structure,
            source_pubchem_cid=source_pubchem_cid,
        )
        if multi_class == "ambiguous_multi_name":
            structure_status = "ambiguous_multi_name"
            docking_ready_reason = "Ambiguous multi-name ingredient record excluded from network and docking."
        if source_pubchem_cid and args.pubchem:
            pubchem_query_status = (
                "matched_structure_by_source_cid"
                if structure_status in {"smiles_ready", "inchikey_and_smiles_ready"}
                else "source_cid_property_lookup_failed"
            )
            pubchem_query_attempted = "true"
        elif source_pubchem_cid:
            pubchem_query_status = "source_cid_not_queried"
            pubchem_query_attempted = "false"
        elif passed and args.pubchem:
            pubchem_query_status = "not_attempted_no_source_cid"
            pubchem_query_attempted = "false"
        else:
            pubchem_query_status = structure.get("pubchem_match_status", "not_applicable")
            pubchem_query_attempted = "false"

        output_row = {
            "herb_name": source_herb_name,
            "ingredient_id": mol_id,
            "raw_ingredient_name": name_info["raw_ingredient_name"],
            "ingredient_name": name,
            "normalized_ingredient_name": name_info["normalized_ingredient_name"],
            "name_repair_applied": name_info["name_repair_applied"],
            "pubchem_cid": structure.get("pubchem_cid", ""),
            "smiles": structure.get("smiles", ""),
            "canonical_smiles": structure.get("canonical_smiles", ""),
            "isomeric_smiles": structure.get("isomeric_smiles", ""),
            "inchikey": structure.get("inchikey", ""),
            "inchi_key": structure.get("inchi_key", ""),
            "pubchem_match_status": structure.get("pubchem_match_status", ""),
            "pubchem_query_name": structure.get("pubchem_query_name", ""),
            "pubchem_attempted_names": structure.get("pubchem_attempted_names", ""),
            "pubchem_error": structure.get("pubchem_error", ""),
            "cid_origin": "source_database_cid" if source_pubchem_cid else "",
            "cid_origin_source_database": source_list_for_row if source_pubchem_cid else "",
            "pubchem_query_attempted": pubchem_query_attempted,
            "pubchem_query_status": pubchem_query_status,
            "smiles_origin": smiles_origin,
            "inchikey_origin": inchikey_origin,
            "structure_file_path": "",
            "structure_status": structure_status,
            "docking_ready_reason": docking_ready_reason,
            "ob": row.get("ob", ""),
            "dl": row.get("dl", ""),
            "source": row.get("source", "") or source_list_for_row,
            "source_list": source_list_for_row,
            "primary_source": row.get("primary_source", "") or "TCMSP",
            "evidence_types": row.get("evidence_types", ""),
            "database_ids": row.get("database_ids", ""),
            "activity_filter_basis": row.get("activity_filter_basis", ""),
            "ingredient_confidence": row.get("ingredient_confidence", ""),
            "source_eligibility": row.get("source_eligibility", "") or row.get("herbiv_category", ""),
            "source_eligibility_reason": row.get("source_eligibility_reason", ""),
            "herbiv_category": row.get("herbiv_category", "") or row.get("source_eligibility", ""),
            "network_eligibility": "eligible" if passed else "excluded",
            "network_eligibility_reason": eligibility_reason,
            "tcmsp_data_source": row_tcmsp_data_source,
            "passed_filter": str(bool(passed)).lower(),
        }
        raw_output.append(output_row)
        if passed:
            active_by_mol[mol_id] = output_row

    if not active_by_mol:
        warnings.append("No network candidate ingredients retained with current source/filter settings; attempting non-structure relaxation while keeping ambiguous multi-name records excluded.")
        for row in raw_output:
            source = next((item for item in ingredient_rows if item.get("MOL_ID") == row["ingredient_id"]), {})
            if row.get("structure_status") != "ambiguous_multi_name" and is_active_ingredient(source, target_mol_ids, filters):
                row["passed_filter"] = "true"
                active_by_mol[row["ingredient_id"]] = row

    target_cache: dict[str, dict[str, str]] = {}
    compound_targets: list[dict] = []
    unmapped_source_rows: list[dict] = []
    ambiguous_source_rows: list[dict] = []
    for row in target_rows:
        mol_id = row.get("MOL_ID", "")
        if mol_id not in active_by_mol:
            continue
        edge_eligible, edge_eligibility_reason = target_edge_eligibility(row)
        if not edge_eligible:
            # The row remains available in union_ingredient_target_edges.csv
            # and source-specific audit files, but cannot enter the main graph.
            continue
        ingredient = active_by_mol[mol_id]
        ingredient_decisions = ingredient_source_eligibility(ingredient, filters)
        final_eligible_sources = target_edge_eligible_sources(row) & {
            source
            for source, (eligible, _reason) in ingredient_decisions.items()
            if eligible
        }
        if not final_eligible_sources:
            # An edge may be provisionally TCMSP-eligible before OB/DL is
            # evaluated. A different source can admit the ingredient without
            # admitting the failed TCMSP evidence on this edge.
            continue
        final_ingredient_confidence, final_relation_confidence = edge_confidence_for_sources(
            row, final_eligible_sources
        )
        target_name = row.get("target_name", "")
        target_gene_hint = str(row.get("target_gene", "") or "").strip()
        if target_gene_hint and target_gene_hint.lower() not in {"none", "nan", "null"}:
            mapping = {
                "target_gene": target_gene_hint.upper(),
                "uniprot_id": row.get("uniprot_id", ""),
                "mapping_status": "mapped_source_gene",
                "mapping_note": (
                    "Gene symbol supplied by the source record; this does not "
                    "independently validate the biological relation."
                ),
            }
        else:
            mapping = map_target_record(target_name, target_cache) if target_name else make_mapping_record("", failure_reason="empty_target_name")
        gene = mapping.get("target_gene", "")
        if not gene:
            issue_row = {
                **mapping,
                "ingredient_id": ingredient["ingredient_id"],
                "raw_ingredient_name": ingredient["raw_ingredient_name"],
                "ingredient_name": ingredient["ingredient_name"],
            }
            if mapping.get("mapping_status") == "ambiguous":
                ambiguous_source_rows.append(issue_row)
            else:
                unmapped_source_rows.append(issue_row)
            continue
        scores = [row.get("SVM_score"), row.get("RF_score")]
        numeric_scores = [value for value in (parse_optional_float(score) for score in scores) if value is not None]
        score = sum(numeric_scores) / len(numeric_scores) if numeric_scores else ""
        raw_evidence_type = str(
            row.get("target_evidence_types")
            or row.get("evidence_type")
            or row.get("validated")
            or ""
        ).strip()
        evidence_class = str(row.get("evidence_class") or "").strip() or classify_target_evidence(
            raw_evidence_type
        )
        mapping_validated = _strict_true(row.get("mapping_validated")) or mapping.get(
            "mapping_status"
        ) == "mapped_uniprot"
        source_priority = {"TCMSP": 1, "HerbiV": 2, "SymMap2": 3}
        eligible_source_list = "; ".join(
            sorted(
                final_eligible_sources,
                key=lambda source: (source_priority.get(source, 99), source),
            )
        )
        all_source_list = str(
            row.get("all_source_list")
            or row.get("source_list")
            or row.get("source")
            or eligible_source_list
        ).strip()
        compound_targets.append(
            {
                "herb_name": source_herb_name,
                "raw_ingredient_name": ingredient.get("raw_ingredient_name", ""),
                "ingredient_name": ingredient["ingredient_name"],
                "normalized_ingredient_name": ingredient.get("normalized_ingredient_name", ""),
                "ingredient_id": ingredient["ingredient_id"],
                "pubchem_cid": ingredient.get("pubchem_cid", ""),
                "smiles": ingredient.get("smiles", ""),
                "canonical_smiles": ingredient.get("canonical_smiles", ""),
                "isomeric_smiles": ingredient.get("isomeric_smiles", ""),
                "inchikey": ingredient.get("inchikey", ""),
                "target_gene": gene,
                "uniprot_id": mapping.get("uniprot_id", ""),
                "target_name": target_name,
                "mapping_status": mapping.get("mapping_status", ""),
                "mapping_note": mapping.get("mapping_note", ""),
                "evidence_type": raw_evidence_type,
                "evidence_class": evidence_class,
                "mapping_validated": str(bool(mapping_validated)).lower(),
                "score": f"{score:.6g}" if isinstance(score, float) else "",
                "source": eligible_source_list,
                "source_list": eligible_source_list,
                "eligible_source_list": eligible_source_list,
                "all_source_list": all_source_list,
                "target_evidence_types": raw_evidence_type,
                "source_target_ids": row.get("source_target_ids", ""),
                "ingredient_confidence": final_ingredient_confidence,
                "relation_confidence": final_relation_confidence,
                "source_eligibility": row.get("source_eligibility", ""),
                "source_eligibility_reason": row.get("source_eligibility_reason", ""),
                "herbiv_category": row.get("herbiv_category", ""),
                "network_eligibility": "eligible",
                "network_eligibility_reason": "; ".join(
                    [
                        edge_eligibility_reason,
                        *[
                            reason
                            for source, (eligible, reason) in ingredient_decisions.items()
                            if source in final_eligible_sources and eligible
                        ],
                    ]
                ),
                "eligibility_status_by_source": row.get("eligibility_status_by_source", ""),
                "eligibility_reason_by_source": row.get("eligibility_reason_by_source", ""),
                "tcmsp_data_source": row_tcmsp_data_source,
            }
        )

    finalize_union_edge_audit(out_root, compound_targets)
    active_rows = [row for row in raw_output if row["passed_filter"] == "true"]
    ligand_ready_rows = build_ligand_ready_rows(active_rows, compound_targets)
    write_module_outputs(
        params=params,
        ingredient_dir=ingredient_dir,
        target_dir=target_dir,
        pubchem_cache_path=pubchem_cache_path,
        pubchem_cache=pubchem_cache,
        shared_pubchem_cache_path=shared_pubchem_cache_path,
        raw_output=raw_output,
        active_rows=active_rows,
        compound_targets=compound_targets,
        unmapped_targets=collect_unmapped_targets(unmapped_source_rows),
        ambiguous_targets=collect_mapping_issues(ambiguous_source_rows),
        ligand_ready_rows=ligand_ready_rows,
        warnings=warnings,
        pubchem_enabled=args.pubchem,
        detail_url=detail_url,
        pubchem_query_attempts=pubchem_query_attempts,
        tcmsp_fetch_diagnostic=tcmsp_diagnostic,
        source_merge_summary=source_merge_summary,
        ambiguous_multi_name_rows=ambiguous_multi_name_rows,
        split_multi_name_rows=split_multi_name_rows,
    )

    print(f"Fetched {len(active_rows)} active ingredients and {len(compound_targets)} compound-target rows.")


if __name__ == "__main__":
    exit_with_errors(main)
